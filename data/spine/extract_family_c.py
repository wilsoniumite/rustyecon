"""Family C extract and validate: land share and factor incomes (Breakpoint B pre-look).

Dated 2026-09-26. Label fetch-3. Fits nothing.

Reads the raw files fetched by fetch_family_c.py (D:/rustyecon-spine/raw/family_c/) and
the hand entries in D:/rustyecon-spine/digitised/family_c/. Writes tidy CSVs to
D:/rustyecon-spine/series/family_c/ and a validation log (validation_family_c.tsv) there.
Nothing it writes goes in the repository: the sources' licences do not allow it (see
manifest_family_c.tsv). Stamp (1916) is public domain; its entries are the exception and
are flagged as redistributable.

Every series carries a construction name. Two sources that cover the same years are
compared and the differences are logged. Nothing is averaged or spliced.

Needs pdftotext (poppler/xpdf) on PATH for the text-layer entries.

Run: D:/rustyecon-spine/venv/Scripts/python.exe data/spine/extract_family_c.py
"""

# %% setup
import os
import re
import subprocess
import sys

import numpy as np
import openpyxl
import pandas as pd

RAW = os.environ.get("SPINE_RAW", "D:/rustyecon-spine/raw/family_c")
DIG = os.environ.get("SPINE_DIG", "D:/rustyecon-spine/digitised/family_c")
OUT = os.environ.get("SPINE_SERIES", "D:/rustyecon-spine/series/family_c")
os.makedirs(OUT, exist_ok=True)

F_NNI = f"{RAW}/clark/England NNI - Clark - 2015.xlsx"
F_C10 = f"{RAW}/clark/Macroagg2009.pdf"
F_C02 = f"{RAW}/clark/rentereh.pdf"
F_BNS10 = f"{RAW}/bns/clark10.xlsx"
F_BOE = f"{RAW}/boe/a-millennium-of-macroeconomic-data-for-the-uk.xlsx"

LOG = []


def log(check, verdict, detail):
    LOG.append(dict(check=check, verdict=verdict, detail=detail))
    print(f"[{verdict:5s}] {check}: {detail}")


def tidy(df, series, unit, geography, source, sheet_column, construction, year_end=None,
         value="value"):
    t = pd.DataFrame({
        "series": series,
        "year": df["year"].astype(int),
        "year_end": (df[year_end] if year_end else df["year"]).astype(int),
        "value": df[value].astype(float),
        "unit": unit,
        "geography": geography,
        "source": source,
        "sheet_column": sheet_column,
        "construction": construction,
    })
    return t


def period_bounds(p):
    """'1860-9' -> (1860, 1869); '1500-39' -> (1500, 1539); '1910-12' -> (1910, 1912)."""
    a, b = p.split("-")
    a = int(a)
    b = int(str(a)[: 4 - len(b)] + b)
    return a, b


def pdftext(path, page):
    r = subprocess.run(["pdftotext", "-raw", "-f", str(page), "-l", str(page), path, "-"],
                       capture_output=True)
    if r.returncode != 0:
        sys.exit(f"pdftotext failed on {path} p{page}: {r.stderr[:200]}")
    return r.stdout.decode("latin-1")


# %% A. Clark (2015) spreadsheet: factor incomes, England, 1209-1869
wb = openpyxl.load_workbook(F_NNI, data_only=True)
ws = wb["Annual"]
cols = dict(year="A", wage_income="I", land_rents="J", net_house_rents="K",
            other_capital_income="L", all_capital_income="M", indirect_taxes="N", nni="O")
rows = []
for r in range(3, 664):  # rows 3-663 are annual; rows 669-735 are decade midpoints
    rows.append({k: ws[f"{c}{r}"].value for k, c in cols.items()})
cl = pd.DataFrame(rows)
ok = cl["year"].tolist() == list(range(1209, 1870))
log("clark2015.annual.coverage", "PASS" if ok else "FAIL",
    f"Annual rows 3-663 hold years {cl.year.min()}-{cl.year.max()}, {len(cl)} rows, contiguous={ok}")
gap = cl.loc[cl.land_rents.isna(), "year"].tolist()
gapw = cl.loc[cl.wage_income.isna(), "year"].tolist()
log("clark2015.annual.gaps", "NOTE",
    f"source gaps, left empty: J-O empty in {len(gap)} years {gap}; I also empty in {sorted(set(gapw) - set(gap))}")
resid = (cl.wage_income + cl.land_rents + cl.all_capital_income + cl.indirect_taxes - cl.nni).abs()
log("clark2015.identity.I+J+M+N=O", "PASS" if resid.max() < 1e-6 * cl.nni.max() else "FAIL",
    f"max |I+J+M+N-O| = {resid.max():.2e} GBP m")
resid2 = (cl.net_house_rents + cl.other_capital_income - cl.all_capital_income).abs()
log("clark2015.identity.K+L=M", "PASS" if resid2.max() < 1e-9 * cl.nni.max() else "FAIL",
    f"max |K+L-M| = {resid2.max():.2e}")

# land rents are not annual observations: find runs of identical values
runs = (cl.land_rents.diff().fillna(1) != 0).cumsum()
rl = cl.groupby(runs).agg(start=("year", "min"), end=("year", "max"))
rl["len"] = rl.end - rl.start + 1
by_era = []
for a, b in [(1209, 1669), (1670, 1841), (1842, 1869)]:
    sub = rl[(rl.start >= a) & (rl.start <= b)]
    by_era.append(f"{a}-{b}: {len(sub)} distinct values, run lengths {sorted(sub['len'].unique().tolist())}")
log("clark2015.land_rents.step_structure", "NOTE", "; ".join(by_era))

cl["land_share_nni"] = cl.land_rents / cl.nni
cl["land_share_factor"] = cl.land_rents / (cl.nni - cl.indirect_taxes)
src15 = "Clark, England NNI - Clark - 2015.xlsx (UC Davis)"
parts = []
for v, colL in [("wage_income", "I"), ("land_rents", "J"), ("net_house_rents", "K"),
                ("other_capital_income", "L"), ("all_capital_income", "M"),
                ("indirect_taxes", "N"), ("nni", "O")]:
    parts.append(tidy(cl.rename(columns={v: "value"}), f"clark2015_{v}", "GBP m, current",
                      "England", src15, f"Annual!{colL}3:{colL}663", "as published"))
pd.concat(parts).to_csv(f"{OUT}/clark2015_factor_incomes_annual.csv", index=False)
pd.concat([
    tidy(cl, "clark2015_land_share_of_nni", "ratio", "England", src15, "Annual!J/O",
         "Clark land rents / Clark NNI (NNI includes indirect taxes)", value="land_share_nni"),
    tidy(cl, "clark2015_land_share_of_factor_income", "ratio", "England", src15,
         "Annual!J/(O-N)", "Clark land rents / (Clark NNI - indirect taxes); the basis of "
         "Clark (2010) Table 13", value="land_share_factor"),
]).to_csv(f"{OUT}/clark2015_land_share_annual.csv", index=False)

# the Decadal sheet against decade means of the Annual rows, and against the midpoint block
wsd = wb["Decadal"]
dec = pd.DataFrame([{k: wsd[f"{c}{r}"].value for k, c in cols.items()} for r in range(3, 70)])
dec = dec.dropna(subset=["year"])
dec["year"] = dec.year.astype(int)
cl["decade"] = (cl.year // 10) * 10
am = cl.groupby("decade")[["land_rents", "nni", "indirect_taxes"]].mean()
cmp_ = dec.set_index("year")[["land_rents", "nni"]].join(am, rsuffix="_annmean")
d1 = (cmp_.land_rents / cmp_.land_rents_annmean - 1).abs()
log("clark2015.decadal_vs_annual_means", "PASS" if d1.max() < 1e-9 else "NOTE",
    f"Decadal!J vs mean of the non-empty Annual!J rows by calendar decade: max rel diff "
    f"{d1.max():.2e} (at {int(d1.idxmax())})")
mid = pd.DataFrame([{k: ws[f"{c}{r}"].value for k, c in cols.items()} for r in range(669, 736)])
mid = mid.dropna(subset=["year"])
same = np.allclose(mid.land_rents.values[: len(dec)], dec.land_rents.values[: len(mid)])
log("clark2015.midpoint_block_eq_decadal", "PASS" if same else "NOTE",
    f"Annual rows 669-735 (decade midpoints, {len(mid)} rows) equal Decadal!J: {same}")
dec["land_share_factor"] = dec.land_rents / (dec.nni - dec.indirect_taxes)
dec["land_share_nni"] = dec.land_rents / dec.nni
dec["year_end"] = dec.year + 9
pd.concat([
    tidy(dec, "clark2015_land_share_of_factor_income_decadal", "ratio", "England", src15,
         "Decadal!J/(O-N)", "Clark land rents / (NNI - indirect taxes), Decadal sheet",
         year_end="year_end", value="land_share_factor"),
    tidy(dec, "clark2015_land_share_of_nni_decadal", "ratio", "England", src15,
         "Decadal!J/O", "Clark land rents / NNI, Decadal sheet", year_end="year_end",
         value="land_share_nni"),
    tidy(dec, "clark2015_land_rents_decadal", "GBP m, current", "England", src15,
         "Decadal!J", "as published", year_end="year_end", value="land_rents"),
    tidy(dec, "clark2015_nni_decadal", "GBP m, current", "England", src15, "Decadal!O",
         "as published", year_end="year_end", value="nni"),
]).to_csv(f"{OUT}/clark2015_land_share_decadal.csv", index=False)

# %% B. Clark (2010) Table 13: text layer (entry A) against the BNS digitisation (entry B)
recs = []
for pg in (65, 66):
    for line in pdftext(F_C10, pg).splitlines():
        m = re.match(r"^(1[2-8]\d0) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)$",
                     line.strip())
        if m:
            recs.append([int(m.group(1))] + [float(x) for x in m.groups()[1:]])
t13a = pd.DataFrame(recs, columns=["year", "nni", "wage_share", "land_share", "capital_share",
                                   "farm_share", "output_per_worker"])
b = pd.read_excel(F_BNS10, sheet_name="clark10_t13")
b.columns = ["year", "nni", "wage_share", "land_share", "capital_share", "farm_share",
             "output_per_worker"]
b = b.dropna(subset=["year"]).astype(float)
b["year"] = b.year.astype(int)
log("clark2010.t13.coverage", "PASS" if len(t13a) == 67 and len(b) == 67 else "FAIL",
    f"text layer {len(t13a)} decades {t13a.year.min()}-{t13a.year.max()}; BNS {len(b)} decades")
j = t13a.set_index("year").join(b.set_index("year"), rsuffix="_bns")
for c in ["nni", "wage_share", "land_share", "capital_share", "farm_share",
          "output_per_worker"]:
    d = (j[c] - j[c + "_bns"]).abs()
    bad = d[d > 1e-9]
    log(f"clark2010.t13.entryA_vs_BNS.{c}", "PASS" if bad.empty else "DIFF",
        "identical in all 67 decades" if bad.empty else
        f"{len(bad)} decades differ: " + ", ".join(f"{k}: pdf {j.loc[k, c]} vs BNS "
                                                   f"{j.loc[k, c + '_bns']}" for k in bad.index[:8]))
s = (t13a.wage_share + t13a.land_share + t13a.capital_share - 1).abs()
log("clark2010.t13.shares_sum_to_1", "PASS" if s.max() <= 0.0015 else "FAIL",
    f"max |wage+land+capital-1| = {s.max():.4f}")
# the farm share and output-per-worker columns repeat themselves 35 rows later
rep = all(np.isclose(t13a.farm_share.values[i], t13a.farm_share.values[i + 35]) and
          np.isclose(t13a.output_per_worker.values[i], t13a.output_per_worker.values[i + 35])
          for i in range(0, 32))
log("clark2010.t13.farm_columns_repeat", "DEFECT" if rep else "PASS",
    "Farm share and output per worker for 1550-1860 repeat the 1200-1510 values row for row "
    "(e.g. 1860 farm share 0.63). A copy error in the working-paper table. The land share "
    "column does not repeat. Do not use those two columns." if rep else "no repetition")
t13a["year_end"] = t13a.year + 9
src10 = "Clark (2010) working paper Macroagg2009.pdf, Table 13 (pp. 65-66)"
pd.concat([
    tidy(t13a, "clark2010_t13_land_share", "ratio", "England", src10, "Table 13 col 4",
         "as published; share of factor income (indirect taxes borne pro rata)",
         year_end="year_end", value="land_share"),
    tidy(t13a, "clark2010_t13_nni", "GBP m, current", "England", src10, "Table 13 col 2",
         "as published", year_end="year_end", value="nni"),
    tidy(t13a, "clark2010_t13_wage_share", "ratio", "England", src10, "Table 13 col 3",
         "as published", year_end="year_end", value="wage_share"),
    tidy(t13a, "clark2010_t13_capital_share", "ratio", "England", src10, "Table 13 col 5",
         "as published", year_end="year_end", value="capital_share"),
]).to_csv(f"{OUT}/clark2010_t13.csv", index=False)

# 2015 spreadsheet against the 2010 table (a vintage comparison, not a transcription check)
v = dec.set_index("year")[["land_share_factor"]].join(t13a.set_index("year")[["land_share"]])
dv = (v.land_share_factor - v.land_share)
w = dv.loc[1750:1860]
log("clark2015_vs_2010.land_share", "NOTE",
    f"2015 J/(O-N) minus 2010 T13 land share: 1200-1860 max |diff| {dv.abs().max():.3f} "
    f"(at {int(dv.abs().idxmax())}); 1750-1860: " +
    ", ".join(f"{k}: {v.loc[k, 'land_share_factor']:.3f} vs {v.loc[k, 'land_share']:.3f}"
              for k in w.index))

# %% C. Clark (2010) Table 34, 1860-2008: text layer (A) against the visual entry (B)
recs = []
for line in pdftext(F_C10, 81).splitlines():
    m = re.match(r"^(\d{4}-\d) (\d+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+) (\d+) ([\d.]+) (\d+)$",
                 line.strip())
    if m:
        recs.append([m.group(1)] + [float(x) for x in m.groups()[1:]])
t34a = pd.DataFrame(recs, columns=["decade", "real_income_pp_index", "share_land_rents",
                                   "share_labor", "share_urban_rents", "share_capital",
                                   "real_wage", "skill_premium", "efficiency_index"])
t34b = pd.read_csv(f"{DIG}/clark2010_t34_entryB_visual.csv")
jj = t34a.set_index("decade").join(t34b.set_index("decade"), rsuffix="_b")
nd = sum(int((jj[c] - jj[c + "_b"]).abs().gt(1e-9).sum()) for c in t34a.columns[1:])
log("clark2010.t34.double_entry", "PASS" if nd == 0 and len(t34a) == 15 else "DIFF",
    f"text layer {len(t34a)} rows vs visual entry {len(t34b)} rows; {nd} cells differ")
s = (t34a.share_land_rents + t34a.share_labor + t34a.share_urban_rents + t34a.share_capital - 1)
log("clark2010.t34.shares_sum_to_1", "PASS" if s.abs().max() <= 0.0015 else "NOTE",
    "as printed: " + ", ".join(f"{d}: {v:+.3f}" for d, v in zip(t34a.decade, s) if abs(v) > 0.0015)
    + " (source rounding or inconsistency; both entries agree)")
t34a[["year", "year_end"]] = pd.DataFrame(t34a.decade.map(period_bounds).tolist())
src34 = "Clark (2010) working paper Macroagg2009.pdf, Table 34 (p. 81)"
pd.concat([
    tidy(t34a, "clark2010_t34_land_share", "ratio", "England", src34, "Table 34 'Share Land Rents'",
         "as published; farmland rents (Stamp 1922 p.49 E&W, adjusted to England, to 1914) over "
         "England NDI (UK NDI scaled by population, Ireland at half weight)",
         year_end="year_end", value="share_land_rents"),
    tidy(t34a, "clark2010_t34_urban_site_rent_share", "ratio", "England", src34,
         "Table 34 'Share Urban Rents'", "as published; Singer (1941) site values to 1910",
         year_end="year_end", value="share_urban_rents"),
    tidy(t34a, "clark2010_t34_labor_share", "ratio", "England", src34, "Table 34 'Share Labor'",
         "as published", year_end="year_end", value="share_labor"),
]).to_csv(f"{OUT}/clark2010_t34.csv", index=False)

# %% D. Clark (2002) second Table 8, net agricultural output: text layer (A) vs visual (B)
recs = []
for line in pdftext(F_C02, 43).splitlines():
    m = re.match(r"^(\d{4}-\d{2}) ([\d.]+) ([\d.]+) (\d+) ([\d.]+) ([\d.]+) ([\d.]+) ([\d.]+)$",
                 line.strip())
    if m:
        recs.append([m.group(1)] + [float(x) for x in m.groups()[1:]])
c8cols = ["period", "population_m", "share_males_agric", "males_agric_000", "farm_wages_gbp_m",
          "land_rents_and_local_taxes_gbp_m", "capital_payments_gbp_m", "net_farm_output_gbp_m"]
t8a = pd.DataFrame(recs, columns=c8cols)
t8b = pd.read_csv(f"{DIG}/clark2002_t8b_entryB_visual.csv", dtype={"period": str})
jj = t8a.set_index("period").join(t8b.set_index("period"), rsuffix="_b")
nd = sum(int((jj[c] - jj[c + "_b"]).abs().gt(1e-9).sum()) for c in c8cols[1:])
log("clark2002.t8b.double_entry", "PASS" if nd == 0 and len(t8a) == 36 else "DIFF",
    f"text layer {len(t8a)} rows vs visual entry {len(t8b)} rows; {nd} cells differ")
s = (t8a.farm_wages_gbp_m + t8a.land_rents_and_local_taxes_gbp_m + t8a.capital_payments_gbp_m -
     t8a.net_farm_output_gbp_m).abs()
# a third entry, if family B has typed the same table (optional; logged, not required)
FB = os.environ.get("SPINE_DIG_B", "D:/rustyecon-spine/digitised/family_b")
for fb in ["clark2002_table8b_entry1_textlayer.csv", "clark2002_table8b_entry2.csv"]:
    pth = f"{FB}/{fb}"
    if os.path.exists(pth):
        o = pd.read_csv(pth, dtype={"period_printed": str}).set_index("period_printed")
        o.columns = c8cols[1:]
        n3 = sum(int((t8a.set_index("period")[c] - o[c]).abs().gt(1e-9).sum()) for c in c8cols[1:])
        log(f"clark2002.t8b.vs_family_b.{fb[:-4]}", "PASS" if n3 == 0 and len(o) == 36 else "DIFF",
            f"family B's entry ({len(o)} rows): {n3} cells differ from ours")
# Table 7's tax burden column (text layer only; a diagnostic, not a series)
tx = {}
for pg in (37, 38):
    for line in pdftext(F_C02, pg).splitlines():
        m = re.match(r"^(\d{4}-\d{2}) [\d.]+ ([\d.]+|-) ", line.strip())
        if m:
            tx[m.group(1)] = np.nan if m.group(2) == "-" else float(m.group(2))
t8a["gap"] = (t8a.net_farm_output_gbp_m - t8a.farm_wages_gbp_m -
              t8a.land_rents_and_local_taxes_gbp_m - t8a.capital_payments_gbp_m)
t8a["gap_over_rent"] = t8a.gap / t8a.land_rents_and_local_taxes_gbp_m
t8a["t7_tax_burden"] = t8a.period.map(tx)
cmpg = t8a[t8a.t7_tax_burden.notna() & (t8a.period >= "1700")]
log("clark2002.t8b.adding_up", "DEFECT" if s.max() > 0.151 else "PASS",
    f"output - (wages + 'rents and local taxes' + capital) is 0.0 to {s.max():.1f} GBP m and grows "
    "over time. The gap over the rent column tracks Table 7's 'tax burden as a share of rents': " +
    ", ".join(f"{r.period} {r.gap_over_rent:.3f} vs {r.t7_tax_burden:.3f}"
              for r in cmpg.itertuples() if r.period[2:4] in ("70", "80", "90", "00", "10", "30", "50"))
    + ". So the column labelled 'Total land rents and local taxes' reads as rent plus tithe "
    "WITHOUT local taxes, and output adds the taxes. Label and arithmetic disagree; not resolved.")
t8a[["year", "year_end"]] = pd.DataFrame(t8a.period.map(period_bounds).tolist())
src02 = "Clark (2002) working paper rentereh.pdf, second Table 8 (p. 43)"
tidy(t8a, "clark2002_t8_land_rents_and_local_taxes", "GBP m, current", "England", src02,
     "Table 8 'Total land rents and local taxes'",
     "as published; farmland rent incl. tithe plus local taxes on occupiers; period means",
     year_end="year_end", value="land_rents_and_local_taxes_gbp_m"
     ).to_csv(f"{OUT}/clark2002_t8_rents.csv", index=False)

# %% E. BoE millennium v3.1: A21 nominal GDP, A17 rent and GDP(I)
bw = openpyxl.load_workbook(F_BOE, read_only=True, data_only=True)
a21 = [list(r) for r in bw["A21. GDP per capita 1086+"].iter_rows(values_only=True)]
a21h = a21[4]
g = pd.DataFrame([r for r in a21[5:] if isinstance(r[0], (int, float))])
g = g.rename(columns={0: "year", 1: "eng_gdp_mp", 3: "eng_gdp_fc", 12: "gb_gdp", 26: "uk_gbi_gdp"})
g = g[["year", "eng_gdp_mp", "eng_gdp_fc", "gb_gdp", "uk_gbi_gdp"]].apply(pd.to_numeric,
                                                                          errors="coerce")
g = g[(g.year >= 1700) & (g.year <= 1920)]
log("boe.a21.headers", "PASS" if (str(a21h[1]).startswith("Nominal GDP at market")
                                  and str(a21h[3]).startswith("Nominal GDP at factor")) else "FAIL",
    f"A21 B='{a21h[1]}', D='{a21h[3]}', M='{a21[4][12]}', AA='{a21[4][26]}'")
srcboe = "Bank of England, A millennium of macroeconomic data v3.1 (file of 2024-09-26)"
pd.concat([
    tidy(g, "boe_a21_england_nominal_gdp_mp", "GBP m, current", "England", srcboe,
         "A21!B", "BoE: GB GDP x England share (Broadberry 1700, Geary-Stark 1861; "
         "Geary-Stark shares of Solomou-Weale UK GDP 1870-1913)", value="eng_gdp_mp"),
    tidy(g, "boe_a21_england_nominal_gdp_fc", "GBP m, current", "England", srcboe, "A21!D",
         "as above, at factor cost", value="eng_gdp_fc"),
    tidy(g, "boe_a21_gb_nominal_gdp", "GBP m, current", "Great Britain", srcboe, "A21!M",
         "as published", value="gb_gdp"),
    tidy(g.dropna(subset=["uk_gbi_gdp"]), "boe_a21_uk_gb_ireland_nominal_gdp", "GBP m, current",
         "UK incl. all Ireland", srcboe, "A21!AA", "as published", value="uk_gbi_gdp"),
]).to_csv(f"{OUT}/boe_a21_nominal_gdp.csv", index=False)

a17 = [list(r) for r in bw["A17. GDP(I) components"].iter_rows(values_only=True)]
from openpyxl.utils import column_index_from_string as CI
hdr = {c: (a17[3][CI(c) - 1], a17[4][CI(c) - 1]) for c in ["AM", "AQ", "AR", "AW"]}
r17 = pd.DataFrame([[r[0], r[CI("AM") - 1], r[CI("AQ") - 1], r[CI("AR") - 1], r[CI("AW") - 1]]
                    for r in a17[5:] if isinstance(r[0], (int, float))],
                   columns=["year", "rent", "dwellings", "gdpi_fc_mitchell", "gdpi_composite"])
r17 = r17[(r17.year >= 1855) & (r17.year <= 1920)].apply(pd.to_numeric)
log("boe.a17.coverage", "PASS" if r17.rent.notna().sum() == 66 else "FAIL",
    f"AM rent 1855-1920 non-empty {r17.rent.notna().sum()}/66; headers {hdr}")
ratio = (r17.dwellings / r17.rent)
log("boe.a17.dwellings_fixed_ratio", "NOTE",
    f"AQ 'capital income from dwellings' / AM 'rent' = {ratio.min():.4f}-{ratio.max():.4f} "
    "1855-1919: a fixed proportion, not an independent series. Rent less dwellings is "
    "therefore not built.")
r17["rent_share_upper"] = r17.rent / r17.gdpi_fc_mitchell
pd.concat([
    tidy(r17, "boe_a17_rent_land_and_buildings", "GBP m, current", "UK", srcboe, "A17!AM",
         "Mitchell (1988) ch. XVI table 5 'Rent' (after Feinstein 1972): land and buildings",
         value="rent"),
    tidy(r17, "boe_a17_gdpi_fc_mitchell", "GBP m, current", "UK", srcboe, "A17!AR",
         "Mitchell (1988) GDP(I) at factor cost", value="gdpi_fc_mitchell"),
    tidy(r17, "boe_a17_rent_share_upper_bound", "ratio", "UK", srcboe, "A17!AM/AR",
         "rent of land AND buildings / GDP(I) factor cost: an upper bound on land's share, "
         "not a land share", value="rent_share_upper"),
]).to_csv(f"{OUT}/boe_a17_rent.csv", index=False)

# %% F. Construction: Clark (2002) rents incl. local taxes / BoE England nominal GDP
gi = g.set_index("year")
out = []
for _, r in t8a.iterrows():
    if r.year < 1700:
        continue
    sl = gi.loc[r.year:r.year_end]
    out.append(dict(year=r.year, year_end=r.year_end, rent=r.land_rents_and_local_taxes_gbp_m,
                    gdp_fc=sl.eng_gdp_fc.mean(), gdp_mp=sl.eng_gdp_mp.mean(), n=len(sl)))
cx = pd.DataFrame(out)
cx["share_fc"] = cx.rent / cx.gdp_fc
cx["share_mp"] = cx.rent / cx.gdp_mp
log("construction.clark2002_over_boe.coverage", "PASS" if (cx.n == cx.year_end - cx.year + 1).all()
    else "FAIL", f"{len(cx)} periods 1700-1912, every year of each period present in A21")
pd.concat([
    tidy(cx, "constr_clark2002_rents_over_boe_england_gdp_fc", "ratio", "England",
         f"{src02}; {srcboe}", "T8 rents / mean(A21!D)",
         "OUR CONSTRUCTION: Clark (2002) farmland rents incl. tithe and local taxes / BoE England "
         "nominal GDP at factor cost (gross of depreciation), period means",
         year_end="year_end", value="share_fc"),
    tidy(cx, "constr_clark2002_rents_over_boe_england_gdp_mp", "ratio", "England",
         f"{src02}; {srcboe}", "T8 rents / mean(A21!B)",
         "OUR CONSTRUCTION: as above over GDP at market prices", year_end="year_end",
         value="share_mp"),
]).to_csv(f"{OUT}/constr_clark2002_over_boe.csv", index=False)

# %% H. Stamp (1916) Table A4: Schedule A 'Lands (including tithes)', 1842-3 to 1913-14
F_STXT = f"{RAW}/stamp/britishincomespr00stamuoft_djvu.txt"
L = open(F_STXT, encoding="utf-8").read().splitlines()
i0 = [i for i, l in enumerate(L) if "including  Tithes)" in l][0]
i1 = [i for i, l in enumerate(L) if "these  tables  see  Appendix" in l and i > i0][0]
yr = re.compile(r"^[1I]\s?[89lI][0-9oOIl]{2}\s*[-\u2013]\s*[0-9oOIlr ]{1,5}")
ocr, cur = [], None
for l in L[i0:i1]:
    t = l.strip()
    if not t:
        continue
    if yr.match(t):
        cur = [re.sub(r"\s", "", yr.match(t).group(0)), []]
        ocr.append(cur)
        continue
    m = re.match(r"^[\dIlOoi][\d,.IlOoi]*", t) if cur else None
    if m:
        v = m.group(0).translate(str.maketrans("IlOoi", "11001")).replace(",", "").replace(".", "")
        if v.isdigit() and int(v) >= 1000:
            cur[1].append(int(v))
# Stamp (1916) is public domain, so its hand entry lives in the repository beside this script
HERE = os.path.dirname(os.path.abspath(__file__))
sb = pd.read_csv(f"{HERE}/stamp1916_a4_lands_entryB_visual.csv", dtype={"fiscal_year": str}, comment="#")
sb = sb[sb.fiscal_year != "1814-15"].reset_index(drop=True)
sb["year"] = sb.fiscal_year.str[:4].astype(int)
ok4 = {int(k[:4]): v for k, v in ((r[0], r[1]) for r in ocr) if len(v) == 4 and k[:4].isdigit()}
unread = sorted(set(sb.year) - set(ok4))
ndiff, diffs = 0, []
for r in sb.itertuples():
    if r.year in ok4:
        for name, a, b_ in zip(["ew", "scotland", "ireland", "uk"], ok4[r.year],
                               [r.ew, r.scotland, r.ireland, r.uk]):
            if a != b_:
                ndiff += 1
                diffs.append(f"{r.fiscal_year} {name}: OCR {a} vs image {b_}")
# cells where OCR and image disagree, each settled by a closer look at the page image
RESOLVED = {"1884-5 scotland": "image reads 7,462 at 300 dpi; E&W+Sc+Ir = UK holds with 7,462",
            "1909-10 uk": "OCR garbled ('52,iir'); image reads 52,111; row sum holds"}
unres = [d for d in diffs if " ".join(d.split()[:2]).rstrip(":") not in RESOLVED]
diffs = [d + (" -> " + RESOLVED[" ".join(d.split()[:2]).rstrip(":")]
              if " ".join(d.split()[:2]).rstrip(":") in RESOLVED else "") for d in diffs]
log("stamp.a4.double_entry", "PASS" if not unres else "DIFF",
    f"OCR text layer parsed cleanly for {len(ok4)} of {len(sb)} years; {ndiff} cells differ "
    f"from the entry read off the page image" + (": " + "; ".join(diffs) if diffs else "") +
    f". OCR unreadable (image entry only, checked by the row sum): {unread}")
rs = sb.ew + sb.scotland + sb.ireland - sb.uk
bad = sb.loc[rs != 0, ["fiscal_year"]].assign(resid=rs[rs != 0])
log("stamp.a4.row_sum", "PASS" if bad.empty else "NOTE",
    f"E&W + Scotland + Ireland = UK holds exactly in {int((rs == 0).sum())} of {len(sb)} years" +
    ("; off in " + ", ".join(f"{r.fiscal_year} ({int(r.resid):+d})" for r in bad.itertuples())
     if not bad.empty else "") + ". 1885-6 and 1890-1 re-read at 300 dpi: printed so. Stamp's "
    "own rounding; the entries stand as printed.")
log("stamp.a4.error_margins", "NOTE",
    "Stamp prints margins (+ or -, GBP 000) for 1842-3 to 1871-2: E&W 100 in 1843-5; "
    "Ireland 300 (1843-52), 100 (1853-61), 20 (1861-72); UK 300-500 (1842-53), 100, 20. "
    "None after 1871-2.")
sb["gb"] = sb.ew + sb.scotland
srcst = "Stamp (1916) British Incomes and Property, Table A4 p.49 (archive.org scan)"
cons = "as published: Sch. A gross assessment, 'Lands' incl. tithe, farmhouses, woodland, building land; fiscal year t/t+1 stored at t"
pd.concat([
    tidy(sb.assign(value=sb.ew / 1000), "stamp_a4_lands_ew", "GBP m, current", "England and Wales",
         srcst, "Table A4 col 2", cons),
    tidy(sb.assign(value=sb.scotland / 1000), "stamp_a4_lands_scotland", "GBP m, current",
         "Scotland", srcst, "Table A4 col 3", cons),
    tidy(sb.assign(value=sb.ireland / 1000), "stamp_a4_lands_ireland", "GBP m, current", "Ireland",
         srcst, "Table A4 col 4", cons),
    tidy(sb.assign(value=sb.uk / 1000), "stamp_a4_lands_uk", "GBP m, current",
         "UK incl. all Ireland", srcst, "Table A4 col 5", cons),
]).to_csv(f"{OUT}/stamp1916_a4_lands.csv", index=False)

# Clark says his land rents from 1842 come from Stamp p.49 (E&W, adjusted to England)
cs = cl.set_index("year").join(sb.set_index("year")[["ew"]], how="inner")
cs["ratio"] = cs.land_rents / (cs.ew / 1000)
j41, j42 = cl.set_index("year").land_rents.loc[[1841, 1842]]
log("clark2015.land_rents.break_1842", "NOTE",
    f"source switch at 1842: 1841 {j41:.2f} -> 1842 {j42:.2f} ({j42 / j41 - 1:+.1%}). Before: "
    "Clark's rent index in quinquennial blocks (1670-1841); from 1842: Sch. A 'Lands' E&W x "
    f"{cs.ratio.mean():.3f} (constant England factor). A level step at the seam, named here.")
log("clark2015_vs_stamp.1842_1869", "NOTE",
    "Clark 2015 land rents / Stamp E&W lands: " + ", ".join(
        f"{y}: {cs.loc[y, 'ratio']:.3f}" for y in [1842, 1845, 1850, 1855, 1860, 1865, 1869]) +
    f"; range {cs.ratio.min():.3f}-{cs.ratio.max():.3f}")

# Construction: Stamp GB lands / BoE GB nominal GDP (same geography, annual)
sg = sb.set_index("year").join(gi[["gb_gdp", "uk_gbi_gdp"]], how="inner")
sg["share_gb"] = (sg.gb / 1000) / sg.gb_gdp
sg["share_uk"] = (sg.uk / 1000) / sg.uk_gbi_gdp
sg = sg.reset_index()
pd.concat([
    tidy(sg, "constr_stamp_lands_gb_over_boe_gb_gdp", "ratio", "Great Britain",
         f"{srcst}; {srcboe}", "(A4 E&W + Scotland) / A21!M",
         "OUR CONSTRUCTION: Sch. A 'Lands' incl. tithe (assessed gross annual value, revalued "
         "periodically) / BoE GB nominal GDP at market prices", value="share_gb"),
    tidy(sg.dropna(subset=["share_uk"]), "constr_stamp_lands_uk_over_boe_uk_gdp", "ratio",
         "UK incl. all Ireland", f"{srcst}; {srcboe}", "A4 UK / A21!AA",
         "OUR CONSTRUCTION: as above, UK incl. all Ireland", value="share_uk"),
]).to_csv(f"{OUT}/constr_stamp_over_boe.csv", index=False)

# %% G. Cross-checks between the land-share constructions (reported, never averaged)
# Clark 2015 rents against Clark 2002 rents (levels)
d15 = dec.set_index("year")
lv = []
for _, r in t8a.iterrows():
    if r.year < 1700 or r.year > 1860:
        continue
    lv.append(f"{r.period}: 2002 {r.land_rents_and_local_taxes_gbp_m:.1f} vs 2015 "
              f"{d15.loc[r.year, 'land_rents']:.1f}")
log("clark2002_vs_clark2015.rent_levels", "NOTE", "; ".join(lv))
# BoE England GDP against Clark NNI, 1700-1869
cm = cl.set_index("year").join(gi, how="inner")
rr = (cm.eng_gdp_fc / (cm.nni - cm.indirect_taxes))
log("boe_gdp_vs_clark_nni", "NOTE",
    "BoE England GDP at factor cost / Clark (NNI - indirect taxes): " +
    ", ".join(f"{y}: {rr.loc[y]:.2f}" for y in [1700, 1750, 1800, 1820, 1850, 1860, 1869]))
# the seam in the 1860s and the post-1869 path, construction by construction
def at(df, y, col):
    return float(df.loc[df.year == y, col].iloc[0])
seam = [
    f"Clark 2015 J/(O-N) 1860s: {at(dec, 1860, 'land_share_factor'):.3f}",
    f"Clark 2015 J/O 1860s: {at(dec, 1860, 'land_share_nni'):.3f}",
    f"Clark 2010 T13 1860: {at(t13a, 1860, 'land_share'):.3f}",
    f"Clark 2010 T34 1860-9: {at(t34a, 1860, 'share_land_rents'):.3f}",
    f"Clark 2002/BoE fc 1860-69: {at(cx, 1860, 'share_fc'):.3f}",
    f"Stamp GB lands/BoE GB GDP 1860-69: {sg[(sg.year>=1860)&(sg.year<=1869)].share_gb.mean():.3f}",
    f"BoE A17 rent(land+buildings)/GDP(I) 1860-69: "
    f"{r17[(r17.year>=1860)&(r17.year<=1869)].rent_share_upper.mean():.3f}",
]
log("seam.1860s", "NOTE", "; ".join(seam))
post = []
for y in [1870, 1880, 1890, 1900, 1910]:
    post.append(f"{y}s: T34 {at(t34a, y, 'share_land_rents'):.3f}, Clark2002/BoE fc "
                f"{at(cx, y, 'share_fc'):.3f}, Stamp GB/BoE GB "
                f"{sg[(sg.year>=y)&(sg.year<=y+9)].share_gb.mean():.3f}, A17 upper "
                f"{r17[(r17.year>=y)&(r17.year<=y+9)].rent_share_upper.mean():.3f}")
log("post1869.constructions", "NOTE", "; ".join(post))

# %% write the validation log
pd.DataFrame(LOG).to_csv(f"{OUT}/validation_family_c.tsv", sep="\t", index=False)
print(f"wrote {OUT}")
