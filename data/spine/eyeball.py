"""The Breakpoint B pre-look: the eyeball sheet's figures and read-off tables.

Dated 2026-09-26. Branch spine-eyeball. Nothing here is fitted: no regression, no
break test, no smoothing, no splice that a source did not make itself.

Reads the tidy series that the three fetch families wrote under $SPINE_SERIES
(default $SPINE_ROOT/series), and one hand entry, Clark (2010) Table 34, from
$SPINE_DIGITISED (default $SPINE_ROOT/digitised), for its real-wage column. Writes:
  - five PNG figures to docs/spine/figs/ (in the repository);
  - read-off tables to $SPINE_EYEBALL (default $SPINE_ROOT/eyeball), outside the
    repository, because most inputs (Bank of England, Clark) may not be redistributed.
    docs/spine/EYEBALL.md quotes these tables.
One check runs here too: the Bank's copy of Allen's wages and CPI against the GPIH copy
of Allen's spreadsheet (fetched by fetch_sheet.py into $SPINE_RAW_SHEET, default
$SPINE_ROOT/raw/sheet, pinned by sha256). It is skipped, and says so, if that file
is absent or its hash differs.

The spine's cache: $SPINE_ROOT, else data/spine/.cache/ beside this script, which git ignores.

Every construction is named where it is built. "Index, base = 100" means each series
is divided by its own mean over the base years: a rescale, not a splice.

Run from the repository root, in the spine's venv (docs/spine/DATA_NOTES.md):
    python data/spine/eyeball.py
"""

# %% setup
import hashlib
import textwrap
import os
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.colors import LinearSegmentedColormap
from matplotlib.ticker import FuncFormatter, NullFormatter

# The spine's cache: $SPINE_ROOT, else data/spine/.cache/ beside this script, which git ignores.
SPINE = Path(os.environ.get("SPINE_ROOT") or Path(__file__).resolve().parent / ".cache")
SER = Path(os.environ.get("SPINE_SERIES", SPINE / "series"))
EYE = Path(os.environ.get("SPINE_EYEBALL", SPINE / "eyeball"))
RAW_SHEET = Path(os.environ.get("SPINE_RAW_SHEET", SPINE / "raw" / "sheet"))
FIGS = Path(__file__).resolve().parents[2] / "docs" / "spine" / "figs"
FIGS.mkdir(parents=True, exist_ok=True)
EYE.mkdir(parents=True, exist_ok=True)

# Validated reference palette (dataviz skill). Three categorical slots per panel at most:
# lines cross, so every pair counts, and only the first three slots pass all pairs.
C1, C2, C3 = "#2a78d6", "#eb6834", "#1baf7a"
INK, INK2, MUTED, GRID, SURF = "#0b0b0b", "#52514e", "#8a8984", "#e4e3df", "#fcfcfb"
BLUE_RAMP = LinearSegmentedColormap.from_list("blue", ["#b7d3f6", "#6da7ec", "#2a78d6",
                                                       "#1c5cab", "#0d366b"])
DATE = "2026-09-26"
plt.rcParams.update({"font.size": 9, "axes.titlesize": 10, "axes.titleweight": "normal",
                     "axes.edgecolor": GRID, "axes.labelcolor": INK2, "xtick.color": INK2,
                     "ytick.color": INK2, "axes.facecolor": SURF, "figure.facecolor": SURF,
                     "savefig.facecolor": SURF, "legend.frameon": False})
READ = []  # read-off lines, written to EYE / readoff.txt


def say(*a):
    s = " ".join(str(x) for x in a)
    print(s)
    READ.append(s)


def fa(name):
    """A family A series (Bank of England carrier) by file stem: year -> value."""
    d = pd.read_csv(SER / "family_a" / "boe" / f"{name}.csv")
    return d.set_index("year")["value"].astype(float).sort_index()


def fb(name):
    d = pd.read_csv(SER / "family_b" / f"{name}.csv")
    return d.set_index("year")["value"].astype(float).sort_index()


def fc(file, series):
    """A family C series: rows with year, year_end, value."""
    d = pd.read_csv(SER / "family_c" / file)
    return d[d.series == series].sort_values("year").reset_index(drop=True)


def dmean(s):
    """Decade means of annual values (decade labelled by its first year)."""
    s = s.dropna()
    return s.groupby((s.index // 10) * 10).mean()


def index(s, a, b):
    """Rescale: divide by the series' own mean over years a..b (inclusive), x 100."""
    base = s.loc[a:b].mean()
    return 100 * s / base


def style(ax, logy=False):
    ax.grid(axis="y", color=GRID, lw=0.7)
    ax.set_axisbelow(True)
    for sp in ["top", "right"]:
        ax.spines[sp].set_visible(False)
    if logy:
        ax.set_yscale("log")
        ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}"))
        ax.yaxis.set_minor_formatter(NullFormatter())


def steps(ax, starts, ends, vals, color, lw=2.2, label=None, **kw):
    """Period values drawn as flat segments over their own years."""
    for i, (a, b, v) in enumerate(zip(starts, ends, vals)):
        ax.plot([a, b + 1], [v, v], color=color, lw=lw, solid_capstyle="butt",
                label=label if i == 0 else None, **kw)


def caption(fig, text, y=0.005):
    width = int(fig.get_figwidth() * 72 / (7 * 0.56))  # characters per line at 7 pt
    fig.text(0.012, y, textwrap.fill(text, width), fontsize=7, color=INK2, va="bottom")


def sha256(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# %% check: the Bank's copy of Allen against the GPIH copy of Allen's spreadsheet
GPIH_ALLEN = RAW_SHEET / "gpih" / "Allen_London_South_Eng_1259-1914.xlsx"
GPIH_PIN = "a8d1b729f86a73498193aea526f1da1d2b82c8098c6d566a6636cf8fe2256827"
if GPIH_ALLEN.exists() and sha256(GPIH_ALLEN) == GPIH_PIN:
    import openpyxl
    from openpyxl.utils import column_index_from_string as ci

    wb = openpyxl.load_workbook(GPIH_ALLEN, read_only=True, data_only=True)
    W = pd.DataFrame(list(wb["Wages"].iter_rows(values_only=True)))
    P = pd.DataFrame(list(wb["Prices"].iter_rows(values_only=True)))

    def gcol(D, j):
        d = D.iloc[8:, [0, j]]
        d = d[pd.to_numeric(d[0], errors="coerce").notna()]
        return pd.Series(pd.to_numeric(d[j], errors="coerce").values,
                         index=d[0].astype(int)).dropna()

    for boe, (D, j, what) in {
        "boe_A47_AG_allen_s_craftsman": (W, 2, "Wages C, S. England craftsman"),
        "boe_A47_AH_allen_s_labourer": (W, 3, "Wages D, S. England building labourer"),
        "boe_A47_AI_allen_s_agric": (W, 4, "Wages E, S. England agricultural labourer"),
        "boe_A47_AJ_allen_london_craftsman": (W, 5, "Wages F, London craftsman"),
        "boe_A47_AK_allen_london_labourer": (W, 6, "Wages G, London building labourer"),
        "boe_A47_BH_allen_cpi_respectable": (P, ci("DQ") - 1, "Prices DQ, NEW CPI (pence)"),
    }.items():
        g, b = gcol(D, j), fa(boe)
        common = g.index.intersection(b.index)
        rel = ((b[common] - g[common]).abs() / g[common].abs()).max()
        same_span = (g.index.min(), g.index.max(), len(g)) == (b.index.min(), b.index.max(), len(b))
        say(f"CHECK allen-carrier {boe} vs GPIH {what}: {len(common)} common years, "
            f"max rel diff {rel:.1e}, same span {same_span}")
    ah, ai = fa("boe_A47_AH_allen_s_labourer"), fa("boe_A47_AI_allen_s_agric")
    common = ah.index.intersection(ai.index)
    eq = ah[common] == ai[common]
    say("CHECK allen agric == building labourer, share of years by half-century:",
        {int(k): round(v, 2) for k, v in eq.groupby((eq.index // 50) * 50).mean().items()})
else:
    say("CHECK allen-carrier SKIPPED: GPIH copy absent or its sha256 differs from the pin")

# %% data: population (C1)
pop_bb = dmean(fa("boe_A2_B_pop_england")).loc[:1860]  # Broadberry et al. 2015 via BoE
pop_cl = fb("clark15_pop_england_decade")  # Clark 2015 spreadsheet, decadal sheet
bench = pd.read_csv(SER / "family_b" / "bns_broadberry15_pop_benchmarks.csv").set_index("year")["value"]

# %% data: real wages, decade means (C1)
w_ann = {
    "clark_avg": fb("cons_clark15_avgwage_over_col_annual"),
    "clark_farm": fb("clark14_real_farm_annual"),  # 1256 printed as 0: dropped by family B
    "clark_lab": fb("clark14_real_bldglab_annual"),
    "clark_craft": fb("clark14_real_craft_annual"),
    "hw16": fa("boe_A48_E_hw2016_real_earnings"),
    "allen_slab": (fa("boe_A47_AH_allen_s_labourer") / fa("boe_A47_BH_allen_cpi_respectable")).dropna(),
    "allen_scraft": (fa("boe_A47_AG_allen_s_craftsman") / fa("boe_A47_BH_allen_cpi_respectable")).dropna(),
    "allen_llab": (fa("boe_A47_AK_allen_london_labourer") / fa("boe_A47_BH_allen_cpi_respectable")).dropna(),
    "boe_comp": fa("boe_A48_B_real_earnings_composite"),
}
wd = pd.DataFrame({k: dmean(v) for k, v in w_ann.items()})
wd["hw19_annual"] = fb("bns_hw19_real_income_annual_contracts_decade")
wd["hw19_day"] = fb("bns_hw19_real_income_day_wages_decade")
# HW (2019) value an annual contract's in-kind benefits at the basket cost, so col F is
# 1 + cash/basket by construction (extract_family_b.py, check V25b). The cash share is the
# only part that moves; it is shown beside col F, never in place of it. Our arithmetic.
wd["hw19_cash_share"] = fb("cons_hw19_cash_over_basket_decade")
wd_full = wd.copy()  # to the end of each series (Allen to the 1910s, BoE composite to 2016)
wd = wd.loc[1200:1860]

NAME = {
    "clark_avg": "Clark average male wage / Clark cost of living",
    "clark_farm": "Clark real farm-labourer day wage",
    "clark_lab": "Clark real building-labourer day wage",
    "clark_craft": "Clark real building-craftsman day wage",
    "hw19_annual": "Humphries-Weisdorf 2019, annual-contract real income",
    "hw19_day": "Humphries-Weisdorf 2019, day-wage real income",
    "hw16": "Humphries-Weisdorf 2016 (BoE A48 E), annual real earnings",
    "allen_slab": "Allen inputs: S. England building labourer / respectable basket",
    "allen_scraft": "Allen inputs: S. England craftsman / respectable basket",
    "allen_llab": "Allen inputs: London building labourer / respectable basket",
    "boe_comp": "BoE composite real consumption earnings (the Bank's splice)",
    "hw19_cash_share": "Humphries-Weisdorf 2019, annual contracts: cash pay / basket cost (our arithmetic)",
}
SHORT = {
    "clark_avg": "Clark avg wage / COL", "clark_farm": "Clark farm", "clark_lab": "Clark bldg labourer",
    "clark_craft": "Clark craftsman", "hw19_annual": "HW 2019 annual contracts",
    "hw19_day": "HW 2019 day wages", "hw16": "HW 2016 (BoE A48 E)",
    "allen_slab": "Allen S. Eng labourer / basket", "allen_scraft": "Allen S. Eng craftsman / basket",
    "allen_llab": "Allen London labourer / basket", "boe_comp": "BoE composite",
    "hw19_cash_share": "HW 2019 cash / basket",
}
wi = pd.DataFrame({k: index(wd[k], 1700, 1749) for k in wd})  # 1700-1749 = 100

# %% read-off: the floor era, decade by decade
say("\n== C1 decade table: population (m) and real wages (index, 1700-49 decades = 100)")
cols = ["clark_avg", "clark_farm", "clark_lab", "clark_craft", "hw19_annual", "hw19_day", "hw16",
        "allen_slab"]
tab = pd.concat([pop_bb.rename("pop_BoE_A2B"), pop_cl.rename("pop_Clark15"), wi[cols],
                 wd["hw19_cash_share"].rename("hw19_cash_share_raw")], axis=1)
tab = tab.loc[1250:1860]
with pd.option_context("display.width", 250, "display.max_rows", 200):
    say(tab.round(2).to_string())
tab.round(4).to_csv(EYE / "c1_decades.csv")

# Every episode EYEBALL.md section 1 quotes. Percent change between decade means (simple,
# not log). HW 2019 cash/basket is shown beside HW 2019 annual (see wd above).
EPISODES = [
    (1260, 1310, "the thirteenth-century rise (from the 1260s, where HW starts)"),
    (1340, 1350, "the Black Death decade"),
    (1340, 1380, "Black Death to the 1380s"),
    (1340, 1450, "Black Death to the population trough"),
    (1500, 1580, "the first eight decades of the sixteenth-century rise"),
    (1500, 1640, "the sixteenth-century rise"),
    (1650, 1740, "the flat century"),
    (1740, 1810, "the late-eighteenth-century rise"),
    (1810, 1840, "the escape decades, to the end of HW"),
    (1810, 1860, "the escape decades, to the end of Clark"),
]
ep_rows = []


def episode(a, b, label):
    say(f"\n-- episode {label}: {a}s -> {b}s")
    for k, s in [("pop_BoE_A2B", pop_bb), ("pop_Clark15", pop_cl)] + [(c, wd[c]) for c in cols + ["hw19_cash_share"]]:
        if a in s.index and b in s.index and pd.notna(s[a]) and pd.notna(s[b]):
            pct = 100 * (s[b] / s[a] - 1)
            say(f"   {k:16s} {s[a]:9.3f} -> {s[b]:9.3f}  ({pct:+.0f}%)")
            ep_rows.append(dict(episode=f"{a}s-{b}s", series=k, start=s[a], end=s[b], pct=round(pct, 2)))


for a, b, label in EPISODES:
    episode(a, b, label)
pd.DataFrame(ep_rows).to_csv(EYE / "c1_episodes.csv", index=False)

say("\n-- population: Broadberry (BoE A2 B) over Clark 2015, 1250-1540")
pop_bb_ann, pop_cl_ann = fa("boe_A2_B_pop_england"), fb("clark15_pop_england_annual")
r_ann = (pop_bb_ann / pop_cl_ann).dropna().loc[1250:1540]
r_dec = (pop_bb / pop_cl).dropna().loc[1250:1530]
say(f"   annual: {r_ann.min():.3f} ({r_ann.idxmin()}) to {r_ann.max():.3f} ({r_ann.idxmax()}); "
    f"decade means: {r_dec.min():.3f} ({r_dec.idxmin()}s) to {r_dec.max():.3f} ({r_dec.idxmax()}s)")
g = 100 * (pop_cl / pop_bb - 1)
say(f"   Clark over Broadberry, decade means: 1600s-1750s {g.loc[1600:1750].min():+.1f}% to "
    f"{g.loc[1600:1750].max():+.1f}%; 1760s-1860s {g.loc[1760:1860].min():+.1f}% to {g.loc[1760:1860].max():+.1f}%")
say("-- Clark's population before 1540 is his 2007 estimate (Clark 2010, section 8). BNS digitisation of Clark "
    "(2007) Table 9: 'best' = the mean of a sample-community series and a series implied by the MPL (the wage).")
c07 = fb("bns_clark07_t9_pop_best_decade")
rr = (pop_cl / c07).dropna().loc[1250:1490]
say(f"   Clark 2015 decadal / Clark 2007 'best', 1250s-1490s: {rr.min():.2f} to {rr.max():.2f}")

say("\n-- 1860s level against the fifteenth-century peak decade (1440s-1450s mean)")
for c in cols + ["boe_comp"]:
    s = wd[c]
    peak = s.loc[1440:1450].mean()
    if pd.notna(s.get(1860)):
        say(f"   {c:12s} 1860s / 1440-50s = {s[1860] / peak:.2f}")
    else:
        last = s.dropna().index.max()
        say(f"   {c:12s} {last}s / 1440-50s = {s[last] / peak:.2f} (series ends {last}s)")

# %% figure 1: population and the real wage on one time axis
fig, axs = plt.subplots(4, 1, figsize=(10, 13.5), sharex=True,
                        gridspec_kw=dict(height_ratios=[1, 1.1, 1.1, 1.1], hspace=0.28))
for ax in axs:
    style(ax, logy=True)
    ax.axvspan(1348, 1352, color=GRID, lw=0)
ax = axs[0]
ax.plot(pop_bb.index + 5, pop_bb.values, color=C1, lw=2, marker="o", ms=3.2,
        label="Broadberry et al. 2015 (BoE A2 col B), decade means")
ax.plot(bench.index, bench.values, ls="none", marker="o", ms=5, mfc=SURF, mec=C1, mew=1.4,
        label="  its benchmark years (the only points before 1541 that are not interpolation)")
ax.plot(pop_cl.index + 5, pop_cl.values, color=C2, lw=2, marker="s", ms=3.2,
        label="Clark 2015 spreadsheet, decadal sheet (1590s repeat the 1580s: flagged)")
ax.axvline(1541, color=MUTED, lw=0.8)
ax.text(1543, 2.0, "1541: BoE series annual\nfrom here (Wrigley line)", fontsize=7.5, color=INK2)
ax.set_ylabel("population of England, millions (log)")
ax.set_title("A. Population of England, 1250-1869", loc="left")
ax.legend(loc="upper left", fontsize=7.5)
ax.set_yticks([2, 3, 4, 6, 10, 15, 20])

panels = [
    (axs[1], "B. Clark's day wages, England (decade means of annual values)",
     [("clark_avg", C1), ("clark_farm", C2), ("clark_lab", C3)]),
    (axs[2], "C. Humphries-Weisdorf constructions, England (annual contracts, day wages x 250, and HW 2016)",
     [("hw19_annual", C1), ("hw19_day", C2), ("hw16", C3)]),
    (axs[3], "D. Allen's inputs via the BoE file (our ratio: day wage / daily cost of the respectable basket)",
     [("allen_slab", C1), ("allen_scraft", C2), ("allen_llab", C3)]),
]
for ax, title, ser in panels:
    for k, col in ser:
        s = wi[k].dropna()
        ax.plot(s.index + 5, s.values, color=col, lw=2, marker="o", ms=3, label=NAME[k])
    ax.axhline(100, color=MUTED, lw=0.7)
    ax.set_title(title, loc="left")
    ax.set_ylabel("index, 1700-1749 = 100 (log)")
    ax.set_ylim(35, 260)
    ax.set_yticks([40, 60, 80, 100, 150, 200, 250])
    ax.legend(loc="upper left", fontsize=7.5, ncol=1)
axs[0].text(1349, 6.2, "Black Death\n1348-51", fontsize=7.5, color=INK2)
axs[-1].set_xlim(1250, 1870)
axs[-1].set_xticks(range(1250, 1871, 50))
axs[-1].set_xlabel("decade (points at mid-decade)")
fig.suptitle("Population and the real wage, England 1250-1869: every reachable construction, unspliced",
             x=0.012, ha="left", y=0.995, fontsize=11.5, color=INK)
caption(fig,
        "Sources. Population: Broadberry, Campbell, Klein, Overton and van Leeuwen (2015) as carried by the Bank of "
        "England millennium dataset v3.1, A2 col B (annual 1541 on; before 1541 log-linear between the 16 benchmarks "
        "drawn hollow); Clark (2015 spreadsheet, 'England NNI - Clark - 2015.xlsx'; before 1540 this is Clark's 2007 "
        "estimate, partly inferred from the real wage through the marginal product of labour, so it is not independent "
        "of the wage panels). Wages: Clark 'Wages 2014.xlsx' real "
        "day wages (Clark's own deflation; the farm wage's 1256 zero dropped); Clark 2015 male average wage / Clark 2015 "
        "cost of living (our ratio); Humphries and Weisdorf (2019) Table A2 cols F and I (BNS digitisation, CC0; col F "
        "values in-kind benefits at the basket cost, so it is 1 + cash/basket by construction); "
        "Humphries and Weisdorf (2016) as BoE A48 col E; Allen's day wages (BoE A47 AG, AH, AK) over Allen's respectable-"
        "basket cost (A47 BH), identical to the GPIH copy of Allen's file. Each wage series divided by its own mean over "
        f"1700-1749 (a rescale). Decade means; no smoothing, no splice. Grey band: 1348-51. {DATE}.", y=0.004)
fig.subplots_adjust(left=0.08, right=0.98, top=0.965, bottom=0.1)
fig.savefig(FIGS / "fig1_floor_time.png", dpi=150)
plt.close(fig)

# %% figures 2 and 3: the phase plots (wage against population, decades dated)
PHASE = ["clark_avg", "clark_farm", "clark_lab", "clark_craft", "hw19_annual", "hw19_day", "hw16",
         "allen_slab"]
LABEL_DECADES = {1250, 1340, 1350, 1450, 1640, 1740, 1800, 1860}


def phase(pop, pop_name, fname, hollow_before=None, pop_note=""):
    fig, axs = plt.subplots(2, 4, figsize=(15, 8.4), sharex=True)
    for ax, k in zip(axs.flat, PHASE):
        style(ax, logy=True)
        ax.set_xscale("log")
        ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}"))
        ax.xaxis.set_minor_formatter(NullFormatter())
        d = pd.concat([pop.rename("N"), wi[k].rename("w")], axis=1).dropna().loc[1250:1860]
        ax.plot(d.N, d.w, color=MUTED, lw=0.8, zorder=1)
        for yr, r in d.iterrows():
            col = C2 if yr >= 1810 else BLUE_RAMP((yr - 1250) / (1800 - 1250))
            hollow = hollow_before is not None and yr < hollow_before
            ax.scatter(r.N, r.w, s=26, zorder=3, color=SURF if hollow else col,
                       edgecolors=col, linewidths=1.3)
            if yr in LABEL_DECADES:
                ax.annotate(f"{yr}s", (r.N, r.w), xytext=(4, 3), textcoords="offset points",
                            fontsize=7, color=INK)
        ax.set_title(SHORT[k], loc="left", fontsize=9.5)
        ax.set_xticks([2, 3, 4, 6, 10, 15, 20])
        ax.set_xlim(1.75, 26)
        ax.set_ylim(35, 260)
        ax.set_yticks([40, 60, 100, 150, 250])
    for ax in axs[1]:
        ax.set_xlabel("population of England, millions (log)")
    for ax in axs[:, 0]:
        ax.set_ylabel("real wage, index 1700-1749 = 100 (log)")
    fig.suptitle(f"The real wage against population, England, decades 1250s-1860s. "
                 f"Population: {pop_name}", x=0.01, ha="left", fontsize=11.5, color=INK)
    note = ("Blue points: 1250s to 1800s, light to dark by date. Orange points: 1810s-1860s. "
            "The line joins decades in time order. ")
    if hollow_before:
        note += (f"Hollow points: decades before {hollow_before}s, where the population decade mean averages "
                 "values interpolated between benchmark years. ")
    caption(fig, note + pop_note + "Wage series and their sources as in figure 1; each divided by its own 1700-1749 "
            f"mean. Nothing fitted: no slope, no line of best fit. {DATE}.", y=0.01)
    fig.subplots_adjust(left=0.05, right=0.99, top=0.92, bottom=0.14, wspace=0.18, hspace=0.25)
    fig.savefig(FIGS / fname, dpi=150)
    plt.close(fig)


phase(pop_bb, "Broadberry et al. 2015 (BoE A2 col B), decade means", "fig2_floor_phase_broadberry.png",
      hollow_before=1540,
      pop_note="Broadberry's population is built without wages: the independent population source here. ")
phase(pop_cl, "Clark 2015 spreadsheet, decadal sheet", "fig3_floor_phase_clark.png",
      pop_note=("Caution: before 1540 Clark's population is his 2007 estimate, half built from the real wage "
                "(the population implied by the marginal product of labour, averaged with sample communities). "
                "Its points before the 1540s are not independent of the wage axis, and the tighter loop they draw is "
                "partly that construction. From 1540 it follows Wrigley et al. (1997), then the census. Its 1590s "
                "repeat the 1580s (3.554 m; flagged). "))

# %% data: the escape (C2), annual, 1700-2016
esc = {
    "boe_comp": fa("boe_A48_B_real_earnings_composite").loc[1700:],
    "clark_avg": w_ann["clark_avg"].loc[1700:],
    "allen_slab": w_ann["allen_slab"].loc[1700:],
    "fein_allenp": fa("boe_A48_X_feinstein_wage_allen_prices"),
    "fein_feinp": fa("boe_A48_W_feinstein_real"),
    "clark_farm": w_ann["clark_farm"].loc[1700:],
}
esc_i = {k: index(v, 1770, 1779) for k, v in esc.items()}
hw = fb("bns_hw19_real_income_annual_contracts_decade").loc[1700:]
hw_i = 100 * hw / hw.loc[1770]
say("\n== C2 decade means, index 1770-1779 = 100")
e_tab = pd.DataFrame({k: dmean(v) for k, v in esc_i.items()})
e_tab["hw19_annual"] = hw_i
with pd.option_context("display.width", 250, "display.max_rows", 200):
    say(e_tab.loc[1700:2010].round(1).to_string())
e_tab.round(3).to_csv(EYE / "c2_decades.csv")
ratio = (esc["boe_comp"] / esc["fein_feinp"]).dropna().loc[1770:1881]
say(f"CHECK BoE composite (A48 B) / Feinstein 1998 real wage (A48 W), annual 1770-1881: relative "
    f"spread of the ratio {ratio.max() / ratio.min() - 1:.1e} over {len(ratio)} years "
    f"(a constant ratio: the composite IS Feinstein's growth there)")
dw = e_tab.loc[1700:1810, ["clark_avg", "clark_farm", "allen_slab"]]
say(f"Day-wage decade means 1700s-1810s (Clark avg/COL, Clark farm, Allen labourer), index 1770s=100: "
    f"{dw.min().min():.1f} to {dw.max().max():.1f}. BoE composite (not a day wage from 1750): "
    f"{e_tab.loc[1700:1810, 'boe_comp'].min():.1f} to {e_tab.loc[1700:1810, 'boe_comp'].max():.1f}")
floor_band = dmean(fa("boe_A48_B_real_earnings_composite")).loc[1250:1740]
floor_band_i = 100 * floor_band / dmean(esc["boe_comp"]).loc[1770]
say(f"BoE composite decade means 1250s-1740s, index 1770s=100: min {floor_band_i.min():.1f} "
    f"({floor_band_i.idxmin()}s), max {floor_band_i.max():.1f} ({floor_band_i.idxmax()}s)")
bc = dmean(esc_i["boe_comp"])
above = bc[bc > floor_band_i.max()]
say(f"BoE composite: first decade above its 1250s-1740s maximum: {above.index.min()}s "
    f"({above.iloc[0]:.1f})")
for k in ["clark_avg", "allen_slab", "clark_farm"]:
    s = dmean(w_ann[k])
    pk = s.loc[1250:1749].max()
    ab = s[(s.index >= 1750) & (s > pk)]
    say(f"{k}: 1250s-1740s max decade {s.loc[1250:1749].idxmax()}s; first decade after 1750 above it: "
        f"{str(ab.index.min()) + 's' if len(ab) else 'none within the series'} (series ends {s.index.max()}s)")
say("Last decade at or below the 1770s level (a read-off, not a break test):")
for k in ["boe_comp", "clark_avg", "allen_slab", "fein_allenp", "clark_farm"]:
    s = dmean(esc_i[k]).loc[1770:]
    lo = s[s <= 100.0001]
    say(f"   {k:12s} {lo.index.max()}s")
lo = hw_i.loc[1770:][hw_i.loc[1770:] <= 100.0001]
say(f"   hw19_annual  {lo.index.max()}s (the 1790s dip to {hw_i[1790]:.1f})")
hwa, hwc = wd_full["hw19_annual"], wd_full["hw19_cash_share"]
say(f"HW19 annual (col F), 1740s-1770s: " + ", ".join(f"{d}s {hwa[d]:.2f}" for d in range(1740, 1780, 10)))
say(f"HW19 1840s over its 1440s-1450s mean: col F {hwa[1840] / hwa.loc[1440:1450].mean():.2f}, cash share "
    f"{hwc[1840] / hwc.loc[1440:1450].mean():.2f}")
bcd = dmean(esc["boe_comp"])
say(f"BoE composite growth per decade, 1810s-1860s: {100 * ((bcd[1860] / bcd[1810]) ** (1 / 5) - 1):.1f}%; "
    f"2000s over 1770s: {bcd[2000] / bcd[1770]:.1f}")
# Clark's own real wage past 1869: Table 34's column, from our visual entry of the table
# (outside the repository; double-entry checked in extract_family_c.py). Not plotted.
DIG = Path(os.environ.get("SPINE_DIGITISED", SPINE / "digitised"))
t34w = pd.read_csv(DIG / "family_c" / "clark2010_t34_entryB_visual.csv").set_index("decade")["real_wage"]
say(f"Clark 2010 Table 34 real wage (Clark 2005 building wages over Feinstein's UK cost of living, then the RPI; "
    f"1860s = 100): " + ", ".join(f"{d} {v:g}" for d, v in t34w.items() if d[:4] in ("1860", "1910", "1950", "2000")))

# %% figure 4: the escape, 1700-2016, log scale
fig, ax = plt.subplots(figsize=(11, 6.6))
style(ax, logy=True)
ax.axhspan(floor_band_i.min(), floor_band_i.max(), color="#eeedea", lw=0, zorder=0)
ax.text(1702, floor_band_i.max() * 1.03, "range of the BoE composite's decade means, 1250s-1740s",
        fontsize=7.5, color=INK2)
ax.axvspan(1810, 1830, color="#fbe3d6", lw=0, zorder=0)
ax.text(1833, 420, "shaded 1810-1830: by eye, the day-wage\nseries turn up for good here",
        fontsize=8, color=INK, va="top")
ax.plot(esc_i["boe_comp"].index, esc_i["boe_comp"].values, color=INK, lw=1.3,
        label="BoE composite real consumption earnings (the Bank's splice; = Feinstein 1998 in 1770-1881)")
ax.plot(esc_i["clark_avg"].index, esc_i["clark_avg"].values, color=C1, lw=1.3,
        label="Clark average male wage / Clark cost of living, England, to 1869")
ax.plot(esc_i["allen_slab"].index, esc_i["allen_slab"].values, color=C2, lw=1.3,
        label="Allen inputs: S. England building labourer / respectable basket, to 1913 (our ratio)")
ax.plot(esc_i["fein_allenp"].index, esc_i["fein_allenp"].values, color=INK2, lw=1.3, ls=(0, (4, 2)),
        label="Feinstein 1998 GB earnings / Allen 2007 prices (BoE A48 X), 1770-1869")
steps(ax, hw_i.index, hw_i.index + 9, hw_i.values, C3, lw=2.4,
      label="Humphries-Weisdorf 2019 annual-contract real income, decades to the 1840s")
ax.axhline(100, color=MUTED, lw=0.7)
ax.set_xlim(1700, 2017)
ax.set_ylim(55, 1800)
ax.set_yticks([60, 80, 100, 150, 200, 300, 500, 1000, 1500])
ax.set_xticks(range(1700, 2001, 25))
ax.set_ylabel("real wage, index 1770-1779 = 100 (log)")
ax.set_title("The escape: the real wage, 1700-2016, five constructions, unspliced "
             "(annual values; HW by decade)", loc="left", fontsize=11, color=INK)
ax.legend(loc="upper left", fontsize=7.6)
caption(fig,
        "Sources. Bank of England millennium dataset v3.1: A48 col B (chain-linked by the Bank: Clark to 1750, "
        "Crafts-Mills 1750-70, Feinstein 1998 1770-1881, Feinstein 1990 to 1911, then Feinstein UK and ONS; England "
        "before 1750, GB then UK after), A48 col X, A47 cols AH and BH (Allen, identical to the GPIH copy). Clark 2015 "
        "spreadsheet cols G/S (our ratio). Humphries and Weisdorf (2019) Table A2 col F (BNS digitisation, CC0). Each "
        "series divided by its own 1770-1779 mean (HW: its 1770s value). The shaded band is a reading by eye, not a "
        f"fitted break. Geography changes at the seams: England, then GB, then UK. {DATE}.")
fig.subplots_adjust(left=0.07, right=0.985, top=0.93, bottom=0.15)
fig.savefig(FIGS / "fig4_escape.png", dpi=150)
plt.close(fig)

# %% data: land's exit (C3)
c15a = fc("clark2015_land_share_annual.csv", "clark2015_land_share_of_factor_income")
c15d = fc("clark2015_land_share_decadal.csv", "clark2015_land_share_of_factor_income_decadal")
c15n = fc("clark2015_land_share_decadal.csv", "clark2015_land_share_of_nni_decadal")
t13 = fc("clark2010_t13.csv", "clark2010_t13_land_share")
t34 = fc("clark2010_t34.csv", "clark2010_t34_land_share")
c02 = fc("constr_clark2002_over_boe.csv", "constr_clark2002_rents_over_boe_england_gdp_fc")
stp = fc("constr_stamp_over_boe.csv", "constr_stamp_lands_gb_over_boe_gb_gdp")
a17 = fc("boe_a17_rent.csv", "boe_a17_rent_share_upper_bound")
say("\n== C3 land share by decade (read off the family C and B series)")
lt = pd.DataFrame({
    "clark15_J/(O-N)": c15d.set_index("year").value, "clark15_J/O": c15n.set_index("year").value,
    "clark10_T13": t13.set_index("year").value, "clark10_T34": t34.set_index("year").value,
    "clark02/BoE_fc": c02.set_index("year").value,
    "stamp_GB/BoE_GB": stp.assign(d=(stp.year // 10) * 10).groupby("d").value.mean(),
    "A17_upper_UK": a17.assign(d=(a17.year // 10) * 10).groupby("d").value.mean(),
    # Clark's 2023 vintage (a private file deposited by BNS): a cross-check column only
    "clark23_NDP_xcheck": dmean(fb("clark23_share_land_ndp_annual")),
})
with pd.option_context("display.width", 250, "display.max_rows", 200):
    say(lt.loc[1250:2000].round(3).to_string())
lt.round(4).to_csv(EYE / "c3_decades.csv")
fl = lt["clark15_J/(O-N)"].loc[1250:1749]
say(f"Clark 2015 J/(O-N), decades 1250s-1740s: min {fl.min():.3f} ({fl.idxmin()}s), "
    f"max {fl.max():.3f} ({fl.idxmax()}s), median {fl.median():.3f}")
f13 = lt["clark10_T13"].loc[1250:1749]
say(f"Clark 2010 T13, decades 1250s-1740s: min {f13.min():.3f}, max {f13.max():.3f}")
for name in ["clark15_J/(O-N)", "clark10_T13", "clark02/BoE_fc"]:
    s = lt[name].dropna().loc[1700:1860]
    below = s[s < fl.min()]
    say(f"{name}: first decade below the 1250s-1740s minimum of Clark 2015 ({fl.min():.3f}): "
        f"{below.index.min() if len(below) else 'none'}s")
    d = s.loc[1750:].diff().dropna()
    say(f"   {name}: decades after the 1750s in which the share rose on the decade before: "
        f"{[int(k) for k, v in d.items() if v > 0]}")
r15, n15 = fb("clark15_land_rents_decade"), fb("clark15_nni_decade")
say(f"Clark 2015, 1770s -> 1860s: land rents x{r15[1860] / r15[1770]:.1f}, NNI x{n15[1860] / n15[1770]:.1f}")

# q-proxy: rent per acre in days of the labourer's wage (our construction)
rent = pd.read_csv(SER / "family_b" / "clark02_t8a_rent_per_acre_england.csv")
wl = fb("gpih_clark06_nominal_bldglab_annual")  # = Wages 2014 to 1869 in every year (V09)
wf = fb("clark14_nominal_farm_annual")
qp = []
for r in rent.itertuples():
    a, b = map(int, r.period.split("-"))
    lab, farm = wl.loc[a:b].mean(), (wf.loc[a:b].mean() if b <= 1869 else np.nan)
    qp.append((a, b, r.value, 240 * r.value / lab, 240 * r.value / farm if farm == farm else np.nan))
qp = pd.DataFrame(qp, columns=["a", "b", "rent_gbp_acre", "days_lab", "days_farm"])
say("\n== C3 proxy: Clark 2002 England rent and tithe per acre, in days of the Clark day wage (our construction)")
say(qp.round(2).to_string())
qp.to_csv(EYE / "c3_days_per_acre.csv", index=False)
q = qp.set_index("a")
for col in ["days_lab", "days_farm"]:
    say(f"   {col}: 1770-74 {q.loc[1770, col]:.2f} -> 1865-69 {q.loc[1865, col]:.2f} "
        f"({100 * (q.loc[1865, col] / q.loc[1770, col] - 1):+.0f}%)")
say(f"   days_lab: 1865-69 {q.loc[1865, 'days_lab']:.2f} -> 1870-74 {q.loc[1870, 'days_lab']:.2f} "
    f"({100 * (q.loc[1870, 'days_lab'] / q.loc[1865, 'days_lab'] - 1):+.0f}%) -> 1895-99 "
    f"{q.loc[1895, 'days_lab']:.2f}; rent per acre 1870-74 -> 1895-99 "
    f"{100 * (q.loc[1895, 'rent_gbp_acre'] / q.loc[1870, 'rent_gbp_acre'] - 1):+.0f}%, building labourer's wage "
    f"{100 * (wl.loc[1895:1899].mean() / wl.loc[1870:1874].mean() - 1):+.0f}%")
say(f"   1865-69 -> 1870-74: building labourer's wage (period mean) "
    f"{100 * (wl.loc[1870:1874].mean() / wl.loc[1865:1869].mean() - 1):+.0f}%, rent per acre "
    f"{100 * (q.loc[1870, 'rent_gbp_acre'] / q.loc[1865, 'rent_gbp_acre'] - 1):+.0f}%")
for y in (1869, 1914):
    say(f"   building labourer's nominal wage {y - 1} -> {y}: {wl[y - 1]:.2f} -> {wl[y]:.2f} d, "
        f"{100 * (wl[y] / wl[y - 1] - 1):+.1f}% (log difference {np.log(wl[y] / wl[y - 1]):+.3f})")
say(f"   rent per acre 1580-99 -> 1600-09: {q.loc[1580, 'rent_gbp_acre']:.3f} -> {q.loc[1600, 'rent_gbp_acre']:.3f}; "
    f"Clark 2015 land rents 1590s -> 1600s: {fb('clark15_land_rents_decade')[1590]:.3f} -> "
    f"{fb('clark15_land_rents_decade')[1600]:.3f} GBP m")

# How independent the land-share constructions are: Clark's rents from 1842 are Stamp's.
stamp_ew = fc("stamp1916_a4_lands.csv", "stamp_a4_lands_ew").set_index("year").value
c15_rent = fb("clark15_land_rents_annual")
rs = (c15_rent / stamp_ew).dropna()
say(f"CHECK Clark 2015 land rents / Stamp Sch. A Lands E&W, {rs.index.min()}-{rs.index.max()} ({len(rs)} years): "
    f"{rs.min():.4f} to {rs.max():.4f}")

# Clark's 2023 vintage (a private file; cross-check only), decade means
c23 = dmean(fb("clark23_share_land_ndp_annual"))
f23 = c23.loc[1250:1740]
say(f"Clark 2023 land share of NDP, decades 1250s-1740s: {f23.min():.3f} ({f23.idxmin()}s) to {f23.max():.3f} "
    f"({f23.idxmax()}s); 1770s {c23[1770]:.3f}, 1840s {c23[1840]:.3f}, 1860s {c23[1860]:.3f}")

# %% read-off: one rule for both sides (section 4). Not a break test and not D2.
# The rule: a series' own range over the decades 1250s-1740s is its floor-era range.
# "Exit" = the first decade after the 1740s outside that range (land below it, wage above
# it), and the first decade of the run outside it that lasts to the series' end. Series
# that start after the 1250s use the decades they have; series that start after 1700
# have no floor-era range and are left out.
say("\n== One rule for both sides: exit from the series' own 1250s-1740s decade range")
exit_rows = []


def exit_decade(s, side, name):
    s = s.dropna()
    fl = s.loc[1250:1749]
    after = s.loc[1750:]
    out = after < fl.min() if side == "land" else after > fl.max()
    first = out[out].index.min() if out.any() else None
    run = None
    if out.any() and out.iloc[-1]:
        k = len(out) - 1
        while k > 0 and out.iloc[k - 1]:
            k -= 1
        run = out.index[k]
    exit_rows.append(dict(series=name, side=side, floor_lo=fl.min(), floor_hi=fl.max(), first_decade=first,
                          first_decade_first_year=fl.index.min(), unbroken_run_from=run, series_end=s.index.max()))
    edge = fl.idxmin() if side == "land" else fl.idxmax()
    say(f"   {side:5s} {name:22s} range {fl.index.min()}s-1740s {fl.min():.3f}-{fl.max():.3f} "
        f"(the edge crossed: {edge}s); first decade out: "
        f"{str(first) + 's' if first else 'none'}; unbroken run out from: {str(run) + 's' if run else 'none'} "
        f"(series ends {s.index.max()}s)")


for k in ["clark_avg", "clark_farm", "clark_lab", "clark_craft", "allen_slab", "hw19_day", "hw19_annual", "hw16",
          "boe_comp"]:
    exit_decade(wd_full[k] if k != "boe_comp" else dmean(fa("boe_A48_B_real_earnings_composite")), "wage", k)
for name, s in [("clark15_J/(O-N)", lt["clark15_J/(O-N)"]), ("clark15_J/O", lt["clark15_J/O"]),
                ("clark10_T13", lt["clark10_T13"]), ("clark23_NDP (cross-check)", c23)]:
    exit_decade(s, "land", name)
pd.DataFrame(exit_rows).to_csv(EYE / "c4_range_exit.csv", index=False)

# %% figure 5: land's exit
fig = plt.figure(figsize=(11, 14))
gs = fig.add_gridspec(3, 1, height_ratios=[1.25, 1, 0.9], hspace=0.3)
ax = fig.add_subplot(gs[0])
style(ax)
ax.plot(a17.year, a17.value, color=MUTED, lw=1.4, ls=(0, (4, 3)),
        label="BoE A17 rent of land AND buildings / GDP(I), UK: an upper bound, not a land share")
ax.plot(c15a.year, c15a.value, color=C1, lw=1.3,
        label="Clark 2015: land rents / (NNI - indirect taxes), England, annual (rents in 5-year blocks to 1841)")
steps(ax, c02.year, c02.year_end, c02.value, C2, lw=2.4,
      label="Our construction: Clark 2002 Table 8 rents / BoE England GDP at factor cost, by decade")
ax.plot(stp.year, stp.value, color=C3, lw=1.6,
        label="Our construction: Stamp 1916 Sch. A 'Lands' (E&W + Scotland) / BoE GB GDP, annual")
steps(ax, t13.year, t13.year_end, t13.value, INK, lw=1.6,
      label="Clark 2010, working paper Macroagg2009: Table 13 to the 1860s, Table 34 from the 1860s (black bars)")
steps(ax, t34.year, t34.year_end, t34.value, INK, lw=1.6)
ax.axvline(1869.5, color=MUTED, lw=0.8)
ax.text(1871, 0.25, "1869: Clark's England\nspreadsheet ends", fontsize=7.5, color=INK2, va="top")
ax.set_xlim(1700, 1921)
ax.set_ylim(0, 0.27)
ax.set_ylabel("share of income")
ax.set_title("A. Land rent's share of income, 1700-1920: every reachable construction", loc="left")
ax.legend(loc="lower left", fontsize=7.4)

ax = fig.add_subplot(gs[1])
style(ax, logy=True)
steps(ax, c15d.year, c15d.year_end, c15d.value, C1, lw=2.2,
      label="Clark 2015 spreadsheet: land rents / (NNI - indirect taxes), decades 1200s-1860s")
steps(ax, t34.year, t34.year_end, t34.value, INK, lw=2.2,
      label="Clark 2010 (working paper) Table 34: farmland rent share of England NDI, 1860s-2000s "
            "(Stamp to 1914; Clark's projection after)")
ax.axhspan(fl.min(), fl.max(), color="#eeedea", lw=0, zorder=0)
ax.text(1430, fl.min() * 0.84, "grey band: range of the Clark 2015 decades, 1250s-1740s", fontsize=7.5,
        color=INK2)
ax.axvspan(1348, 1352, color=GRID, lw=0)
ax.set_xlim(1200, 2010)
ax.set_ylim(0.0015, 0.45)
ax.set_yticks([0.002, 0.005, 0.01, 0.02, 0.05, 0.1, 0.2, 0.3])
ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:g}"))
ax.set_ylabel("share of income (log)")
ax.set_title("B. The same over seven centuries, log scale (England)", loc="left")
ax.legend(loc="lower left", fontsize=7.4)

ax = fig.add_subplot(gs[2])
style(ax, logy=True)
steps(ax, qp.a, qp.b, qp.days_lab, C1, lw=2.2,
      label="rent and tithe of one acre for a year / building labourer's day wage (Clark, GPIH copy), 1500-1912")
steps(ax, qp.a, qp.b, qp.days_farm, C2, lw=2.2,
      label="rent and tithe of one acre for a year / farm labourer's day wage (Clark 2014), 1500-1869")
ax.axvline(1869, color=MUTED, lw=0.8)
ax.text(1866, 4.6, "1869: the building labourer's\nwage steps +14.1% in one year", fontsize=7.5, color=INK2,
        va="top", ha="right")
ax.set_xlim(1495, 1917)
ax.set_ylim(2.5, 20)
ax.set_yticks([3, 4, 5, 7, 10, 15, 20])
ax.set_ylabel("days of wage per acre-year (log)")
ax.set_title("C. A price proxy (our construction): Clark 2002 England rent and tithe per acre, in days of the "
             "day wage", loc="left")
ax.legend(loc="upper left", fontsize=7.4)
fig.suptitle("Land's exit: land rent against income, and the price of an acre against a day's labour",
             x=0.01, ha="left", y=0.995, fontsize=11.5, color=INK)
caption(fig,
        "Sources. Clark 'England NNI - Clark - 2015.xlsx' cols J, N, O (land rents: one value per decade before 1670, per 5-year block 1670-1841; from 1842 "
        "Schedule A 'Lands' x 0.939). Clark (2010) Tables 13 and 34 (working paper 'Macroagg2009.pdf', double entry; "
        "Table 13 also checked against the BNS digitisation). Table 34's rents are Stamp's to 1914; after 1914 Clark "
        "projects them with Feinstein's UK farm rents to 1944, then uses DEFRA land prices x an assumed 3% return x an "
        "assumed 28 m acres (1945-67), then DEFRA tenancy rents x 28 m acres: the 1910s-1920s step falls on that seam. "
        "Not independent: before 1842 every land share here rests on Clark's rents; from 1842 Clark 2015, Tables 13 and "
        "34 and our Stamp construction all rest on Schedule A 'Lands' (Clark 2015's rents = 0.9386 x Stamp E&W; "
        "'Lands' includes tithe, farmhouses, woodland and building land), and Clark 2002 is Clark's own rent data. "
        "Clark (2002) working paper 'rentereh.pdf', the two tables "
        "numbered 8 (double entry; the rents column is labelled 'rents and local taxes', and its arithmetic reads as "
        "rent plus tithe without local taxes: unresolved). Panel C: Clark's first Table 8 is rent and tithe per acre; "
        "its England averages for 1500-39 and 1560-79 have no North figure; rent per acre rises 2.5-fold from 1580-99 "
        "to 1600-09 as printed (Clark 2015's land rents double from the 1590s to the 1600s too). Stamp (1916) British "
        "Incomes and Property, Table A4 "
        "(public domain; double entry). Bank of England millennium dataset v3.1, A17 cols AM/AR (Mitchell 1988), A21 "
        "(England GDP = the Bank's England share x GB GDP). GPIH copy of Clark's wages (2006 vintage, equal to Wages "
        f"2014 through 1869). Period values drawn over their own years. Nothing spliced, nothing smoothed. {DATE}.")
fig.subplots_adjust(left=0.08, right=0.985, top=0.965, bottom=0.125)
fig.savefig(FIGS / "fig5_land_exit.png", dpi=150)
plt.close(fig)

# %% write the read-off
(EYE / "readoff.txt").write_text("\n".join(READ) + "\n", encoding="utf-8")
print("\nwrote", *sorted(p.name for p in FIGS.glob("*.png")), "to", FIGS)
print("wrote read-off tables to", EYE)
