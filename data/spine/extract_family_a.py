# Extract family A's series into tidy CSV: the Bank of England millennium file and the
# BNS replication archive.
#
# Dated 2026-09-26. Branch spine-eyeball. Nothing here is fitted.
#
# Reads the raw files that fetch_family_a.py put under $SPINE_RAW/family_a (default
# D:/rustyecon-spine/raw/family_a). Writes one CSV per series under
# $SPINE_SERIES/family_a (default D:/rustyecon-spine/series/family_a), plus
# series_index.csv and all_long.csv.
#
# Tidy columns: series_id, year, value, unit, geography, source, carrier, sheet, column,
# basis, note.
#   source  = who made the numbers, as the carrier's own header attributes them.
#   carrier = the file we read them from (BoE v3.1 as served 2026-09-26; the BNS deposit).
#   basis   = how this row came to exist: benchmark, interpolated, chain-linked segment,
#             decade mean, and so on. Every construction is named; none is hidden.
#
# The BoE series are copyright the Bank and third parties. They stay out of the public
# repository (manifest: derived_in_repo = no). The BNS outputs are CC0.
#
# Run: D:/rustyecon-spine/venv/Scripts/python.exe data/spine/extract_family_a.py

# %% imports and places
import csv
import math
import os
import pathlib
import re
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
from openpyxl.utils import column_index_from_string as CI

RAW = pathlib.Path(os.environ.get("SPINE_RAW", "D:/rustyecon-spine/raw")) / "family_a"
OUT = pathlib.Path(os.environ.get("SPINE_SERIES", "D:/rustyecon-spine/series")) / "family_a"
BOE = RAW / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
BOE_BNS = RAW / "bns" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
CARRIER_BOE = "BoE, A millennium of macroeconomic data for the UK, v3.1 (file of 2024-09-26, fetched 2026-09-26)"
CARRIER_BOE_BNS = "BoE v3.1, earlier file (modified 2018-09-13) as deposited in BNS doi:10.7910/DVN/5EXFLU"
CARRIER_BNS = "BNS replication archive, Harvard Dataverse doi:10.7910/DVN/5EXFLU v2.2 (CC0)"
for sub in ("boe", "boe_bns_vintage", "bns"):
    (OUT / sub).mkdir(parents=True, exist_ok=True)
COLS = ["series_id", "year", "value", "unit", "geography", "source", "carrier", "sheet", "column", "basis", "note"]

# %% the BoE series: (id, sheet, column, first data row, first year, unit, geography, source, note)
# The first data row and year are where column A starts its run of years in that sheet.
A2, A3, A9, A17 = "A2. Pop of Eng & GB 1086-1870", "A3. Eng. Agriculture 1270-1870", "A9. Nominal GDP (A)", "A17. GDP(I) components"
A18, A21, A32 = "A18. Population 1680+", "A21. GDP per capita 1086+", "A32. Property prices & rent"
A47, A48 = "A47. Wages and prices", "A48. Real Earnings "
START = {A2: (9, 1086), A3: (13, 1270), A9: (6, 1688), A17: (6, 1855), A18: (6, 1086), A21: (6, 1086),
         A32: (5, 1845), A47: (7, 1209), A48: (6, 1209)}
BROAD = "Broadberry, Campbell, Klein, Overton and van Leeuwen (2015), British Economic Growth 1270-1870"
BOE_SERIES = [
    # population
    ("boe_A2_B_pop_england", A2, "B", "millions", "England", BROAD, "Population of England, millions"),
    ("boe_A2_C_pop_gb", A2, "C", "millions", "Great Britain", BROAD, "Population of Great Britain, millions"),
    ("boe_A18_K_pop_england", A18, "K", "thousands", "England", "Wrigley (1997) to 1841, Census thereafter (BoE label)",
     "formula: =Z to 1841, =AA from 1842"),
    ("boe_A18_Z_pop_england", A18, "Z", "thousands", "England", "Wrigley (1997), Broadberry et al. (2015) (BoE label)",
     "formula: =1000 x A2 col B, 1086-1841"),
    ("boe_A18_AA_pop_england_census", A18, "AA", "thousands", "England",
     "Census estimates, incl. armed forces overseas in WW1 (BoE label)", "formula: England and Wales (P) minus Wales (R)"),
    ("boe_A21_G_pop_england", A21, "G", "millions", "England", "BoE A21 (links A2 col B to 1700, A18 col K from 1701)",
     "formula links, not a separate source"),
    # land quantity
    ("boe_A3_B_arable_acreage", A3, "B", "million acres", "England", BROAD + ", Table 2.10",
     "Total arable and sown acreage"),
    # national income denominators
    ("boe_A21_B_gdp_england_nominal_mp", A21, "B", "GBP million, current prices, market prices", "England",
     "BoE A21: to 1699 = A21 C + D; from 1700 = A9 col L", "nominal GDP of England at market prices"),
    ("boe_A21_C_factor_cost_adj", A21, "C", "GBP million", "England",
     "BoE A21: factor cost adjustment 1270-1700 (direct taxes on land treated as a tax on production)", ""),
    ("boe_A21_D_gdp_england_nominal_fc", A21, "D", "GBP million, current prices, factor cost", "England",
     "BoE A21: to 1699 chained back on Broadberry et al. A6 col G nominal GDP; from 1700 = A9 col J", ""),
    ("boe_A21_E_gdp_england_real_mp", A21, "E", "GBP million, 2013 prices, market prices", "England",
     "BoE A21: before 1700 chained on Broadberry et al. A6 col E (real GDP, 1700=100); from 1700 = A8 col P", ""),
    ("boe_A21_F_gdp_england_real_fc", A21, "F", "GBP million, 2013 prices, factor cost", "England",
     "BoE A21: before 1700 chained on Broadberry et al. A6 col E; from 1700 = A8 col N", ""),
    ("boe_A21_H_gdp_pc_england_nominal", A21, "H", "GBP per head, current prices", "England",
     "BoE A21: = col B / col G", "divides by the interpolated population before 1541"),
    ("boe_A21_I_gdp_pc_england_real", A21, "I", "GBP per head, 2013 prices", "England",
     "BoE A21: = col E / col G", "divides by the interpolated population before 1541"),
    ("boe_A9_J_gdp_england_nominal_fc", A9, "J", "GBP million, current prices, factor cost", "England",
     "BoE A9: 1700-1910 = England share of GB (col CU) x GB GDP at factor cost (col BM); later links per A9",
     "a share-of-GB construction before 1911"),
    ("boe_A9_CU_england_share_of_gb", A9, "CU", "percent of GB GDP", "England/GB",
     "Geary and Stark (2015a) and Broadberry et al. (2015), interpolated (BoE label)", ""),
    ("boe_A9_CW_clark_nni", A9, "CW", "GBP million, current prices", "England",
     "Clark (2015) (BoE memo: Net National Income alternative estimate)", "memo column"),
    # rent and income, 1855+
    ("boe_A17_AM_rent_mitchell", A17, "AM", "GBP million, current prices", "UK (GB + Ireland before 1920)",
     "Mitchell (1988), ch. XVI, Table 5, pp. 831-835 (BoE label)", "Rent: land AND buildings; not a land rent"),
    ("boe_A17_AN_rent_mitchell_1920_65", A17, "AN", "GBP million, current prices", "UK",
     "Mitchell (1988), ch. XVI, Table 5, pp. 831-836 (BoE label)", "Rent, second block"),
    ("boe_A17_AR_gdpi_mitchell_fc", A17, "AR", "GBP million, current prices, factor cost", "UK (GB + Ireland before 1920)",
     "Mitchell (1988), ch. XVI, Table 5, pp. 828-830 (BoE label)", "GDP(I) at factor cost; same table as AM"),
    ("boe_A17_AW_gdpi_composite", A17, "AW", "GBP million, current prices", "UK",
     "BoE composite GDP(I): chained on AR to 1919, AU 1920-1947, AV from 1948", "chain-linked; level set by the modern end"),
    ("boe_A17_AQ_dwellings_income", A17, "AQ", "GBP million, current prices", "UK",
     "BoE spliced capital income from dwellings (chained on AM/AN growth before 1948)", ""),
    ("boe_A17_Z_rent_buildings", A17, "Z", "GBP million, current prices", "UK", "ONS DTWR, extrapolated (BoE label)", ""),
    ("boe_A32_AC_agri_land_price", A32, "AC", "GBP thousand per acre", "UK",
     "Vallis (1972), Urban land and building prices 1892-1969, Estates Gazette (BoE label)",
     "a land PRICE, not a rent"),
    ("boe_A32_P_house_rents_index", A32, "P", "index 1900=100", "Britain",
     "Holmans, Historical statistics of housing in Britain, Table H.2 (BoE label)", "house rents, not land"),
    # nominal wages (d/day) and price indices
    ("boe_A47_B_earnings_composite", A47, "B", "GBP per week (chain-linked; level = ONS 2016)", "England/GB (BoE label)",
     "BoE composite average weekly earnings: English daily wage rates before 1750 (BoE note)", "splice: see basis"),
    ("boe_A47_D_cpi_composite", A47, "D", "index 2015=100 (chain-linked)", "GB/UK (England before 1750)",
     "BoE preferred CPI splice (Clark COL, Schumpeter-Gilboy, Crafts-Mills, Feinstein 1998, Feinstein 1991, ONS)",
     "splice: see basis"),
    ("boe_A47_Z_clark_farm_wage", A47, "Z", "pence per day", "England", "Clark (2009) (BoE label)", "farm wages"),
    ("boe_A47_AB_clark_craftsman_wage", A47, "AB", "pence per day", "England", "Clark (2009) (BoE label)", "building craftsmen"),
    ("boe_A47_AC_clark_labourer_wage", A47, "AC", "pence per day", "England", "Clark (2009) (BoE label)", "building labourers"),
    ("boe_A47_AE_clark_avg_wage_levi", A47, "AE", "pence per day", "England",
     "Average wage by the formula in Clark (2009), benchmarked to Levi (1867) (BoE label)", ""),
    ("boe_A47_AF_clark_avg_male_wage", A47, "AF", "pence per day", "England",
     "Average male wage based on Clark (2015), interpolations in red (BoE label)", ""),
    ("boe_A47_AG_allen_s_craftsman", A47, "AG", "pence per day", "Southern England",
     "Allen (2013), daily wage rates downloaded from GPIH (BoE label)", "building craftsman"),
    ("boe_A47_AH_allen_s_labourer", A47, "AH", "pence per day", "Southern England",
     "Allen (2013), daily wage rates downloaded from GPIH (BoE label)", "building labourer"),
    ("boe_A47_AI_allen_s_agric", A47, "AI", "pence per day", "Southern England",
     "Allen (2013), daily wage rates downloaded from GPIH (BoE label)", "agricultural labourer"),
    ("boe_A47_AJ_allen_london_craftsman", A47, "AJ", "pence per day", "London",
     "Allen (2013), daily wage rates downloaded from GPIH (BoE label)", "building craftsman"),
    ("boe_A47_AK_allen_london_labourer", A47, "AK", "pence per day", "London",
     "Allen (2013), daily wage rates downloaded from GPIH (BoE label)", "building labourer"),
    ("boe_A47_AP_craftsmills_earnings", A47, "AP", "index 1900=100", "GB",
     "Crafts and Mills (1994), average weekly wage earnings 1750-1880 (BoE label)", ""),
    ("boe_A47_AQ_feinstein_earnings_gbp", A47, "AQ", "GBP per year", "GB",
     "Feinstein (1998), Appendix Table 1, pp. 652-3: average full-employment money earnings (BoE label)", ""),
    ("boe_A47_AR_feinstein_earnings_index", A47, "AR", "index 1778/82=100", "GB",
     "Feinstein (1998), Appendix Table 1, pp. 652-3 (BoE label)", ""),
    ("boe_A47_BA_feinstein_earnings_uk_gbp", A47, "BA", "GBP per year", "UK (pre-1920)",
     "Feinstein (1998), Appendix Table 1, pp. 652-3 (BoE label)", ""),
    ("boe_A47_BB_feinstein_earnings_uk_index", A47, "BB", "index 1778/82=100", "UK (pre-1920)",
     "Feinstein (1998), Appendix Table 1 (BoE label)", ""),
    ("boe_A47_BG_clark_col", A47, "BG", "index 1860-9=100", "England", "Clark (2009, 2014) cost of living (BoE label)", ""),
    ("boe_A47_BH_allen_cpi_respectable", A47, "BH", "pence per day (cost of the respectable basket)", "Southern England/London",
     "Allen (2013) (BoE label)", ""),
    ("boe_A47_BI_schumpeter_gilboy", A47, "BI", "index 1701=100", "England",
     "Schumpeter-Gilboy consumer price index, Mitchell (1988) p. 719 (BoE label)", ""),
    ("boe_A47_BJ_craftsmills_col", A47, "BJ", "index 1900=100", "GB", "Crafts and Mills (1994) cost of living (BoE label)", ""),
    ("boe_A47_BK_feinstein_col", A47, "BK", "index 1778/82=100", "GB",
     "Feinstein (1998), Appendix Table 1, pp. 652-3, cost of living (BoE label)", ""),
    ("boe_A47_BL_allen2007_col", A47, "BL", "index 1860-69=100", "GB", "Allen (2007), Appendix 1 (BoE label)", ""),
    # real wages
    ("boe_A48_B_real_earnings_composite", A48, "B", "index 1900=100", "GB (BoE label); England daily wages before 1750",
     "BoE composite real consumption earnings = 100 x (A47 B / A47 D), rebased 1900", "a splice: see basis"),
    ("boe_A48_E_hw2016_real_earnings", A48, "E", "ratio as given (annual earnings / annual cost of Allen's basket)", "England",
     "Humphries and Weisdorf (2016): annual male earnings incl. payments in kind BY DECADE / yearly CPI from Allen (2013) incl. a 5% rent adjustment (BoE label)",
     "numerator decadal, denominator annual: the within-decade movement is prices only"),
    ("boe_A48_F_days_adjusted", A48, "F", "ratio (chained to col E at 1750)", "England",
     "BoE construction: daily wage rates adjusted for days worked from a range of estimates (A54 col I)", "BoE's own construction"),
    ("boe_A48_G_real_earnings_days_adjusted", A48, "G", "index 1900=100", "GB/England",
     "BoE construction: = col B from 1760; chained on col F before", "BoE's own construction"),
    ("boe_A48_V_clark_real_earnings", A48, "V", "pence per day / (Clark COL, 1860-9=100)", "England",
     "Clark (2009, 2014) (BoE label); formula = A47 AF / A47 BG", ""),
    ("boe_A48_W_feinstein_real", A48, "W", "index 1835=100", "GB", "Feinstein (1998) (BoE label); = A47 AR / A47 BK", ""),
    ("boe_A48_X_feinstein_wage_allen_prices", A48, "X", "index 1835=100", "GB",
     "Feinstein (1998) wages, Allen (2007) prices (BoE label); = A47 AQ / A47 BL", ""),
    ("boe_A48_Y_feinstein_uk_wage_gb_prices", A48, "Y", "index 1835=100", "UK wages / GB prices",
     "Feinstein (1998) UK wages and GB prices (BoE label); = A47 BB / A47 BK", ""),
    ("boe_A48_Z_feinstein_uk_wage_allen_prices", A48, "Z", "index 1835=100", "UK wages / GB prices",
     "Feinstein (1998) UK wages, Allen (2007) GB prices (BoE label); = A47 BB / A47 BL", ""),
]

# %% how each composite was built (read off the formulas in the file; see validate)
# Links run backwards: the value at year y is the value at y+1 times col(y)/col(y+1).
B_WAGE_LINKS = [(1209, 1749, "AF Clark avg male daily wage, England"), (1750, 1769, "AP Crafts-Mills weekly earnings, GB"),
                (1770, 1880, "AR Feinstein 1998 earnings, GB"), (1881, 1910, "AS Feinstein 1990 earnings, GB"),
                (1911, 1912, "BC Feinstein earnings, UK"), (1913, 1962, "BD Feinstein weekly earnings, UK"),
                (1963, 2016, "BF ONS weekly wages, GB")]
D_PRICE_LINKS = [(1209, 1660, "BG Clark COL, England"), (1661, 1749, "BI Schumpeter-Gilboy, England"),
                 (1750, 1769, "BJ Crafts-Mills COL, GB"), (1770, 1881, "BK Feinstein 1998 COL, GB"),
                 (1882, 1913, "BW Feinstein 1991 COL"), (1914, 2016, "F ONS spliced CPI")]
BROADBERRY_BENCHMARKS = [1086, 1190, 1220, 1250, 1279, 1290, 1315, 1325, 1348, 1351, 1377, 1400, 1430, 1450, 1522, 1541]


def link(table, y):
    for a, b, lab in table:
        if a <= y <= b:
            return lab
    return ""


def basis_for(sid, year, red_years):
    if sid == "boe_A2_B_pop_england":
        if year in BROADBERRY_BENCHMARKS:
            return "Broadberry benchmark year"
        if year < 1250:
            return "interpolated: BoE formula, geometric between benchmarks"
        if year < 1541:
            return "interpolated: log-linear between Broadberry benchmarks"
        return "annual value"
    if sid == "boe_A18_Z_pop_england":
        return "= A2 col B x 1000; " + basis_for("boe_A2_B_pop_england", year, red_years)
    if sid == "boe_A18_K_pop_england":
        return ("= col Z (A2 col B x 1000); " + basis_for("boe_A2_B_pop_england", year, red_years)) if year <= 1841 \
            else "= col AA (census, England and Wales minus Wales)"
    if sid == "boe_A21_G_pop_england":
        return ("= A2 col B; " + basis_for("boe_A2_B_pop_england", year, red_years)) if year <= 1700 \
            else "= A18 col K / 1000; " + basis_for("boe_A18_K_pop_england", year, red_years)
    if sid == "boe_A18_AA_pop_england_census":
        return "England and Wales (col P) minus Wales (col R)"
    if sid in ("boe_A47_AF_clark_avg_male_wage", "boe_A47_BG_clark_col", "boe_A48_V_clark_real_earnings"):
        return "interpolated (red font in the BoE file)" if year in red_years else "as given"
    if sid == "boe_A47_B_earnings_composite":
        return "link to next year: " + link(B_WAGE_LINKS, year)
    if sid == "boe_A47_D_cpi_composite":
        return "link to next year: " + link(D_PRICE_LINKS, year)
    if sid == "boe_A48_B_real_earnings_composite":
        return "wage link: " + link(B_WAGE_LINKS, year) + "; price link: " + link(D_PRICE_LINKS, year)
    if sid == "boe_A48_G_real_earnings_days_adjusted":
        return "= col B" if year >= 1760 else "chained on col F (days-adjusted)"
    if sid == "boe_A48_F_days_adjusted":
        return "= col E (HW 2016)" if year >= 1750 else "BoE days-worked adjustment of daily wages, chained"
    if sid == "boe_A9_J_gdp_england_nominal_fc":
        return "England share of GB x GB GDP" if year <= 1910 else "BoE later links (A9 cols CZ, CY/EE, DD, EE)"
    if sid == "boe_A21_D_gdp_england_nominal_fc":
        return "chained on Broadberry A6 col G" if year < 1700 else "= A9 col J"
    if sid == "boe_A21_B_gdp_england_nominal_mp":
        return "= A21 C + D" if year < 1700 else "= A9 col L"
    if sid in ("boe_A21_E_gdp_england_real_mp", "boe_A21_F_gdp_england_real_fc"):
        return "chained on Broadberry A6 col E" if year < 1700 else "= A8 (UK real GDP sheet), England column"
    if sid in ("boe_A21_H_gdp_pc_england_nominal", "boe_A21_I_gdp_pc_england_real"):
        return "ratio to A21 G; population " + basis_for("boe_A21_G_pop_england", year, red_years)
    if sid == "boe_A17_AW_gdpi_composite":
        return "chained on AR (Mitchell)" if year <= 1919 else ("chained on AU" if year <= 1947 else "= AV (ONS)")
    return "as given"


# %% red-font cells (the Bank marks Clark's interpolations in red)
NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}


def red_font_years(xlsx, sheet_name, col, r0, y0):
    z = zipfile.ZipFile(xlsx)
    wbx = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    rmap = {r.get("Id"): r.get("Target") for r in rels}
    target = [rmap[s.get("{%s}id" % NS["r"])] for s in wbx.find("m:sheets", NS) if s.get("name") == sheet_name][0]
    st = ET.fromstring(z.read("xl/styles.xml"))
    fonts = st.find("m:fonts", NS).findall("m:font", NS)
    xfs = st.find("m:cellXfs", NS).findall("m:xf", NS)
    red = set()
    with z.open("xl/" + target.lstrip("/").replace("xl/", "")) as f:
        for ev, el in ET.iterparse(f):
            if el.tag.endswith("}c"):
                m = re.match(r"([A-Z]+)(\d+)", el.get("r"))
                if m.group(1) == col and int(m.group(2)) >= r0:
                    font = fonts[int(xfs[int(el.get("s", 0))].get("fontId", 0))]
                    c = font.find("m:color", NS)
                    if c is not None and (c.get("rgb") or "").upper().endswith("FF0000") and el.find("m:v", NS) is not None:
                        red.add(y0 + int(m.group(2)) - r0)
                el.clear()
    return red


# %% read a BoE workbook into {sheet: rows}
def read_boe(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for sh in {s[1] for s in BOE_SERIES}:
        r0, _ = START[sh]
        sheets[sh] = list(wb[sh].iter_rows(min_row=r0, values_only=True))
    return sheets


def boe_rows(sheets, sid, sh, col, unit, geo, src, note, carrier, red_years):
    r0, y0 = START[sh]
    ci = CI(col) - 1
    out = []
    for i, row in enumerate(sheets[sh]):
        yr = row[0]
        if not isinstance(yr, (int, float)):
            continue
        assert int(yr) == y0 + i, (sh, yr, y0 + i)  # the year column has no gaps
        v = row[ci] if ci < len(row) else None
        if isinstance(v, (int, float)) and not (isinstance(v, float) and math.isnan(v)):
            out.append(dict(series_id=sid, year=int(yr), value=repr(float(v)), unit=unit, geography=geo, source=src,
                            carrier=carrier, sheet=sh.strip(), column=col, basis=basis_for(sid, int(yr), red_years),
                            note=note))
    return out


def write(rows, path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)


# %% extract BoE, both vintages
RED = red_font_years(BOE, A47, "AF", 7, 1209)
RED_BG = red_font_years(BOE, A47, "BG", 7, 1209)
assert RED == RED_BG, (RED, RED_BG)
print("red-font (interpolated) years in A47 AF and BG:", sorted(RED))
index, all_rows = [], []
for path, carrier, sub in [(BOE, CARRIER_BOE, "boe"), (BOE_BNS, CARRIER_BOE_BNS, "boe_bns_vintage")]:
    sheets = read_boe(path)
    for sid, sh, col, unit, geo, src, note in BOE_SERIES:
        rows = boe_rows(sheets, sid, sh, col, unit, geo, src, note, carrier, RED)
        write(rows, OUT / sub / f"{sid}.csv")
        if sub == "boe":
            all_rows += rows
            index.append(dict(series_id=sid, file=f"boe/{sid}.csv", first=rows[0]["year"] if rows else "",
                              last=rows[-1]["year"] if rows else "", n=len(rows), unit=unit, geography=geo,
                              source=src, carrier=carrier, sheet=sh.strip(), column=col, redistribute="no"))
    print(sub, "done")

# %% BNS: malthus_data.csv (BNS's assembled decadal input; construction per dataset.R)
MD = {
    "pop_10": ("millions", "England", "Clark (2010) population via Steinsson MalthusFigures.xlsx (dataset.R)", ""),
    "w_farm_10": ("index 1860s=100", "England", "Clark, Wages 2014.xlsx Decadal, Real Farm Wage (dataset.R)", ""),
    "w_lab_10": ("index 1860s=100", "England", "Clark, Wages 2014.xlsx Decadal, Real Building Laborer Wage (dataset.R)", ""),
    "w_craft_10": ("index 1860s=100", "England", "Clark, Wages 2014.xlsx Decadal, Real Building Craftsman Wage (dataset.R)", ""),
    "allen_w": ("index, rescaled to w_lab_10 in the 1770s", "GB / England",
                "Allen (2007) wage, decade mean, rescaled to Clark w_lab_10 in 1770; FILLED WITH Clark w_lab_10 where Allen is missing (dataset.R)",
                "a splice: Clark outside 1770-1860"),
    "inc_hw": ("ratio as given", "England", "Humphries-Weisdorf (2019) Table A2, Real income (annual wages) (dataset.R)", ""),
    "d_hw": ("days", "England", "BNS construction: HW (2019) implied income / Clark farm day wage (dataset.R)", ""),
    "broad_pop_hf": ("millions", "England", "Decade mean of BoE A2 col B (BNS's 2018 BoE copy) (dataset.R)", "decade mean of an interpolated annual series"),
    "broad_pop_lf": ("millions", "England", "Broadberry et al. (2015) Table 1.06 benchmarks by decade before 1540, else broad_pop_hf (dataset.R)", ""),
    "wrigley_pop": ("millions", "England", "= pop_10 (Clark) from 1540 (dataset.R): NOT a separate Wrigley series", "named by BNS, not by source"),
    "pop_07": ("millions", "England", "Clark (2007) Table 9, best population estimate (dataset.R)", ""),
    "land_share": ("share of NNI", "England", "Clark (2010) Table 13, Land (dataset.R)", ""),
    "w_ind": ("ratio", "England", "Clark NNI 2015 Decadal: male average wage / price index of domestic expenditure (dataset.R)", ""),
    "rent_ind": ("ratio", "England", "Clark NNI 2015 Decadal: land rents / price index of domestic expenditure (dataset.R)", ""),
}
with open(RAW / "bns" / "malthus_data.csv", encoding="utf-8") as f:
    md = list(csv.DictReader(f))
for col, (unit, geo, src, note) in MD.items():
    rows = [dict(series_id=f"bns_malthus_{col}", year=int(r["decade"]), value=r[col], unit=unit, geography=geo, source=src,
                 carrier=CARRIER_BNS, sheet="malthus_data.csv", column=col, basis="decade (start year)", note=note)
            for r in md if r[col] not in ("", "NA")]
    write(rows, OUT / "bns" / f"bns_malthus_{col}.csv")
    all_rows += rows
    index.append(dict(series_id=f"bns_malthus_{col}", file=f"bns/bns_malthus_{col}.csv", first=rows[0]["year"],
                      last=rows[-1]["year"], n=len(rows), unit=unit, geography=geo, source=src, carrier=CARRIER_BNS,
                      sheet="malthus_data.csv", column=col,
                      redistribute="flag: CC0 as deposited, but a compilation of third-party numbers; not committed"))


# %% BNS digitisations (xlsx)
def xl(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    ws.reset_dimensions()
    return list(ws.iter_rows(values_only=True))


def add(sid, pairs, unit, geo, src, sheet, col, basis, note, redis):
    rows = [dict(series_id=sid, year=int(y), value=repr(float(v)), unit=unit, geography=geo, source=src, carrier=CARRIER_BNS,
                 sheet=sheet, column=col, basis=basis, note=note) for y, v in pairs if isinstance(v, (int, float))]
    write(rows, OUT / "bns" / f"{sid}.csv")
    all_rows.extend(rows)
    index.append(dict(series_id=sid, file=f"bns/{sid}.csv", first=rows[0]["year"], last=rows[-1]["year"], n=len(rows),
                      unit=unit, geography=geo, source=src, carrier=CARRIER_BNS, sheet=sheet, column=col, redistribute=redis))


DIG = "yes (CC0 digitisation); the numbers are the cited authors'"
hw = xl(RAW / "bns" / "hw19.xlsx", "hw19_tabA2")
hwd = [(int(str(r[0])[:4]), r) for r in hw[2:] if r[0] and re.match(r"\d{4}", str(r[0]))]
HWSRC = "Humphries and Weisdorf (2019), EJ 129, Table A2 (BNS digitisation)"
for c, name, unit in [(4, "E_implied_income_annual", "pence per year (implied income, annual contracts)"),
                      (5, "F_real_income_annual", "ratio as given (real income, annual contracts)"),
                      (6, "G_day_pay", "pence per day"), (7, "H_implied_income_day", "pence per year (implied income, day wages)"),
                      (8, "I_real_income_day", "ratio as given (real income, day wages)"),
                      (9, "J_cpi_day_pay_respectable", "ratio (CPI / day pay, respectable)"),
                      (10, "K_cpi_day_pay_barebones", "ratio (CPI / day pay, bare bones)")]:
    add(f"bns_hw19_{name}", [(d, r[c]) for d, r in hwd], unit, "England", HWSRC, "hw19_tabA2",
        "ABCDEFGHIJKL"[c], "decade (start year of the table's decade label)", "", DIG)
a1 = xl(RAW / "bns" / "allen07.xlsx", "allen07_app1")
a2 = xl(RAW / "bns" / "allen07.xlsx", "allen07_app2")
add("bns_allen07_cpi", [(r[0], r[1]) for r in a1[1:]], "index 1860-69=100", "GB",
    "Allen (2007), Pessimism preserved, Appendix I (BNS digitisation)", "allen07_app1", "B", "annual", "", DIG)
add("bns_allen07_wage", [(r[0], r[1]) for r in a2[1:]], "as digitised; the digitisation states no unit", "GB",
    "Allen (2007), Pessimism preserved, Appendix II (BNS digitisation)", "allen07_app2", "B", "annual",
    "unit not stated by BNS; do not read as a level without the paper", DIG)
t13 = xl(RAW / "bns" / "clark10.xlsx", "clark10_t13")
for c, name, unit in [(1, "NNI", "as printed in Table 13 (column NNI)"), (2, "wage_share", "share of NNI"), (3, "land_share", "share of NNI"),
                      (4, "capital_share", "share of NNI"), (5, "farm_share", "share"), (6, "output_per_worker", "as printed")]:
    add(f"bns_clark10_t13_{name}", [(r[0], r[c]) for r in t13[1:]], unit, "England",
        "Clark (2010), Macroeconomic aggregates for England, Table 13 (BNS digitisation)", "clark10_t13", "ABCDEFG"[c],
        "decade (start year)", "", DIG)
t1 = xl(RAW / "bns" / "clark10.xlsx", "clark10_t1")
for c, name in [(1, "farm_wage"), (2, "building_labourer_wage"), (3, "building_craftsman_wage"), (5, "average_wage")]:
    add(f"bns_clark10_t1_{name}", [(r[0], r[c]) for r in t1[1:]], "pence per day (nominal, as printed)", "England",
        "Clark (2010), Table 1 (BNS digitisation)", "clark10_t1", "ABCDEFG"[c], "decade (start year)", "", DIG)
c07 = xl(RAW / "bns" / "clark07.xlsx", "clark07_t9")
add("bns_clark07_pop_best", [(int(str(r[0])[:4]), r[6]) for r in c07[1:] if r[0]], "millions", "England",
    "Clark (2007), The long march of history, Table 9, best estimate from MPL and sample communities (BNS digitisation)",
    "clark07_t9", "G", "decade (start year)", "", DIG)
add("bns_clark07_pop_sample_scaled", [(int(str(r[0])[:4]), r[4]) for r in c07[1:] if r[0]], "millions", "England",
    "Clark (2007), Table 9, sample population scaled to national levels (BNS digitisation)", "clark07_t9", "E",
    "decade (start year)", "", DIG)
bb = xl(RAW / "bns" / "broadberry_etal15.xlsx", "bckol15_tab1.06")
add("bns_broadberry15_benchmarks", [(r[0], r[1]) for r in bb[1:]], "millions", "England",
    BROAD + ", Table 1.06 benchmark estimates (BNS digitisation)", "bckol15_tab1.06", "B", "benchmark year", "", DIG)
est = xl(RAW / "bns" / "bns_estimates.xlsx", "baseline")
h = est[0]
for name in ["N_mean", "N_p5", "N_p95"]:
    c = h.index(name)
    add(f"bns_estimates_baseline_{name}", [(r[0], r[c]) for r in est[1:] if r[0]], "millions", "England",
        "Bouscasse, Nakamura and Steinsson, When did growth begin?, baseline posterior (bns_estimates.xlsx)", "baseline", name,
        "decade; MODEL OUTPUT, not data", "a labelled reference only", "yes (CC0)")

# %% index and long file
write(all_rows, OUT / "all_long.csv")
with open(OUT / "series_index.csv", "w", encoding="utf-8", newline="") as f:
    w = csv.DictWriter(f, fieldnames=list(index[0].keys()), lineterminator="\n")
    w.writeheader()
    w.writerows(index)
print(len(index), "series;", len(all_rows), "rows ->", OUT)
