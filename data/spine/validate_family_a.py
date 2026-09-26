# Validate family A's series: the BoE millennium file and the BNS archive.
#
# Dated 2026-09-26. Branch spine-eyeball. Nothing here is fitted, averaged or spliced.
# Where two sources cover the same years the differences are reported, not reconciled.
#
# Reads the tidy CSVs from extract_family_a.py and, for cross-checks, the raw files:
# the BNS deposit's copies of Clark's spreadsheets (third-party; used only to check what
# the BoE file says it carries from Clark) and the manifest.
# Writes $SPINE_SERIES/family_a/validation_family_a.tsv and prints it.
#
# Result codes: PASS (an identity or equality holds), DIFF (two carriers or sources
# differ; the size is given), FLAG (a construction, splice, interpolation or label the
# user of the series must know), INFO (coverage and descriptive numbers), STOP (a failure
# that blocks use).
#
# Run: D:/rustyecon-spine/venv/Scripts/python.exe data/spine/validate_family_a.py

# %% imports and places
import csv
import math
import os
import pathlib
import re

import openpyxl
import pandas as pd

HERE = pathlib.Path(__file__).resolve().parent
RAW = pathlib.Path(os.environ.get("SPINE_RAW", "D:/rustyecon-spine/raw")) / "family_a"
OUT = pathlib.Path(os.environ.get("SPINE_SERIES", "D:/rustyecon-spine/series")) / "family_a"
BOE = RAW / "boe" / "a-millennium-of-macroeconomic-data-for-the-uk.xlsx"
results = []


def rec(cid, crit, subject, result, detail):
    results.append(dict(check=cid, criterion=crit, subject=subject, result=result, detail=detail))


def S(sid, sub=None):
    """Load one tidy series as a year-indexed float Series."""
    sub = sub or ("boe" if sid.startswith("boe_") else "bns")
    df = pd.read_csv(OUT / sub / f"{sid}.csv")
    return pd.Series(df["value"].astype(float).values, index=df["year"].astype(int).values, name=sid)


def f4(x):
    return "nan" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.4g}"


def maxrel(a, b):
    """Max relative difference over the years where both carry a number; n counts those years."""
    a, b = a.align(b, join="inner")
    keep = a.notna() & b.notna()
    a, b = a[keep], b[keep]
    d = ((a - b).abs() / b.abs().where(b != 0)).dropna()
    return (float(d.max()) if len(d) else float("nan")), len(a), d


idx = pd.read_csv(OUT / "series_index.csv")
BOE_IDS = [s for s in idx["series_id"] if s.startswith("boe_")]

# %% V01: the fetch (manifest)
man = list(csv.DictReader(open(HERE / "manifest_family_a.tsv", encoding="utf-8"), delimiter="\t"))
bad = [m["key"] + ": " + m["status"] for m in man if re.match(r"(BLOCKED|SHA256|MD5|NOT IN)", m["status"])]
rec("V01a", "all", "fetch", "STOP" if bad else "PASS",
    "; ".join(bad) if bad else f"{len(man)} files fetched or cached; every pinned sha256 matched")
dv = [m for m in man if "dataverse" in m["url"]]
rec("V01b", "all", "Dataverse MD5", "PASS" if all(m["md5_check"] == "ok" for m in dv) else "STOP",
    f"{sum(m['md5_check'] == 'ok' for m in dv)}/{len(dv)} files match the MD5 the archive publishes")
sha = {m["key"]: m["sha256"] for m in man}
rec("V01c", "C1", "bns_estimates.xlsx: author site vs CC0 deposit",
    "PASS" if sha["bns_site_estimates"] == sha["bns_estimates"] else "DIFF", "identical sha256" if
    sha["bns_site_estimates"] == sha["bns_estimates"] else "sha256 differ")
boe_m = [m for m in man if m["key"] == "boe_millennium"][0]
rec("V01d", "all", "BoE file", "INFO",
    f"{boe_m['bytes']} B, Last-Modified {boe_m['last_modified']}; sheet 'Corrections to V3.1' lists Aug 2018 and "
    "Aug 2024 corrections (A4 col E wool/textiles 1701-1870; M10/M13 monthly consols and shares). None touches a "
    "column used here. The BNS copy (29,766,578 B, docProps modified 2018-09-13) has 'What's new in V3.1' with the "
    "Aug 2018 list only")

# %% V02: vintage diff, the Bank's file today against the BNS deposit's 2018 copy
ndiff_total = 0
for sid in BOE_IDS:
    a, b = S(sid, "boe"), S(sid, "boe_bns_vintage")
    only_a, only_b = sorted(set(a.index) - set(b.index)), sorted(set(b.index) - set(a.index))
    m, n, d = maxrel(a, b)
    nd = int((d > 1e-12).sum())
    ndiff_total += nd + len(only_a) + len(only_b)
    if nd or only_a or only_b:
        yrs = list(d[d > 1e-12].index[:8])
        rec("V02", "all", sid, "DIFF", f"{nd} of {n} common years differ (max rel {f4(m)}; first {yrs}); "
            f"only in 2024 file: {only_a[:8]}; only in 2018 copy: {only_b[:8]}")
rec("V02", "all", "all BoE columns used", "PASS" if ndiff_total == 0 else "DIFF",
    f"{len(BOE_IDS)} columns compared year by year across the two vintages; {ndiff_total} differing cells")

# %% V03: coverage, gaps and sign
for sid in BOE_IDS:
    s = S(sid)
    span = set(range(s.index.min(), s.index.max() + 1))
    gaps = sorted(span - set(s.index))
    nonpos = int((s <= 0).sum())
    g = ""
    if gaps:
        runs, start, prev = [], gaps[0], gaps[0]
        for y in gaps[1:]:
            if y != prev + 1:
                runs.append((start, prev))
                start = y
            prev = y
        runs.append((start, prev))
        g = f"; {len(gaps)} missing years inside the span, runs {runs[:6]}{' ...' if len(runs) > 6 else ''}"
    rec("V03", "all", sid, "FLAG" if (gaps or nonpos) else "INFO",
        f"{s.index.min()}-{s.index.max()}, n={len(s)}, min {f4(s.min())}, max {f4(s.max())}"
        f"{'; ' + str(nonpos) + ' values <= 0' if nonpos else ''}{g}")

# %% V04: the BoE population columns are one series
b = S("boe_A2_B_pop_england")
z = S("boe_A18_Z_pop_england")
k = S("boe_A18_K_pop_england")
aa = S("boe_A18_AA_pop_england_census")
g = S("boe_A21_G_pop_england")
m1, n1, _ = maxrel(z, b * 1000)
m2, n2, _ = maxrel(k[k.index <= 1841], z)
m3, n3, _ = maxrel(k[k.index >= 1842], aa)
m4, n4, _ = maxrel(g[g.index <= 1700], b)
m5, n5, _ = maxrel(g[g.index >= 1701], k / 1000)
ok = max(m1, m2, m3, m4, m5) < 1e-12
rec("V04", "C1", "BoE population: A2 B, A18 Z, A18 K, A21 G", "FLAG" if ok else "DIFF",
    f"A18 Z = 1000 x A2 B ({n1} yrs, max rel {f4(m1)}); A18 K = Z to 1841 ({n2}), = AA census from 1842 ({n3}); "
    f"A21 G = A2 B to 1700 ({n4}), = K/1000 from 1701 ({n5}). These are ONE series (Broadberry et al. to 1841, "
    "census after), linked by formulas, not three sources. The label 'Wrigley (1997)' on A18 K and Z is not a "
    "second source here: the cells are A2 col B x 1000")

# %% V05-V06: Broadberry's benchmarks and the interpolation between them
bb = S("bns_broadberry15_benchmarks")
diffs = {int(y): b[y] - v for y, v in bb.items()}
mx = max(abs(v) for v in diffs.values())
rec("V05", "C1", "A2 B at Broadberry's Table 1.06 benchmark years vs the BNS digitisation", "PASS" if mx <= 0.005 else "DIFF",
    f"16 benchmarks 1086-1541; max |diff| {mx:.4f} million (BoE unrounded, table at 2 decimals): "
    + ", ".join(f"{y}:{v:+.4f}" for y, v in diffs.items() if abs(v) > 1e-9))
lb = b.apply(math.log)
kinks = [y for y in range(1087, 1541) if abs(lb[y + 1] - 2 * lb[y] + lb[y - 1]) > 1e-9]
rec("V06a", "C1", "A2 B, 1086-1540: where the log path bends", "FLAG",
    f"bends only at {kinks}. Between these years the series is log-linear: interpolation. 1087-1249 are BoE formula "
    "cells (geometric). 1250-1540 are typed values that are still log-linear between benchmarks. 1348->1351 is a "
    "three-step geometric fall (4.814, 3.922, 3.195, 2.603). Only the benchmark years are data before 1541")
post = [y for y in range(1542, 1870) if abs(lb[y + 1] - 2 * lb[y] + lb[y - 1]) > 1e-9]
rec("V06b", "C1", "A2 B, 1541-1870", "INFO",
    f"the log path bends in {len(post)} of 328 years: annual values, not interpolation between quinquennia")
rec("V06c", "C1", "annual population around 1348", "FLAG",
    "no annual population exists in family A's sources for 1250-1540; BoE annual values there are interpolation. "
    "An annual Black Death window needs an annual population source; none found (PLAN Breakpoint A question 2)")

# %% V07: the census seam in A18 K (1841/1842)
m7, n7, _ = maxrel(z[z.index >= 1801], aa[aa.index <= 1841])
rec("V07", "C1", "A18: Broadberry/Wrigley (Z) against the census column (AA) in their overlap 1801-1841",
    "PASS" if m7 < 1e-6 else "DIFF",
    f"{n7} years, max rel diff {f4(m7)}: the census column reproduces Z in the overlap, so K has no level step at "
    f"1842 (K1841={k[1841]:.0f}, K1842={k[1842]:.0f}, {100 * (k[1842] / k[1841] - 1):+.2f}% vs "
    f"{100 * (k[1841] / k[1840] - 1):+.2f}% the year before). AA is England and Wales (P) minus Wales (R)")

# %% V08: BoE population against Clark's (two sources, same years)
wbc = openpyxl.load_workbook(RAW / "bns" / "England NNI - Clark - 2015.xlsx", read_only=True, data_only=True)
ws = wbc["Annual"]
ws.reset_dimensions()
crow = list(ws.iter_rows(min_row=3, max_row=663, values_only=True))
assert crow[0][0] == 1209 and crow[-1][0] == 1869
cy = [int(r[0]) for r in crow]
def cell(r, c):
    i = ord(c) - 65
    return r[i] if i < len(r) and isinstance(r[i], (int, float)) else float("nan")


col = {c: pd.Series([cell(r, c) for r in crow], index=cy) for c in "CGJOS"}
dec = [1250, 1300, 1340, 1350, 1360, 1380, 1400, 1450, 1500, 1540, 1600, 1650, 1700, 1750, 1800, 1850, 1860]
rows = [f"{y}:{b[y] / col['C'][y]:.3f}" for y in dec if y in b.index and not math.isnan(col["C"][y])]
lr = (b.reindex(col["C"].index).apply(math.log) - col["C"].apply(math.log)).dropna()
rec("V08", "C1", "BoE A2 B (Broadberry) / Clark 2015 'Pop England' (BNS copy of Clark), same years", "DIFF",
    "ratio at " + ", ".join(rows) + f". Largest gap {lr.abs().idxmax()} ({100 * (math.exp(lr[lr.abs().idxmax()]) - 1):+.0f}%). "
    "Two sources; report both, never average")
pop10 = S("bns_malthus_pop_10")
pd_ = {y: pop10[y] / col["C"][y] for y in [1210, 1300, 1400, 1500, 1590, 1600, 1700, 1800, 1860] if y in pop10.index and y in col["C"].index}
rec("V08b", "C1", "BNS pop_10 (Clark via Steinsson) vs Clark 2015 annual 'Pop England' at the decade's first year", "INFO",
    ", ".join(f"{y}:{v:.3f}" for y, v in pd_.items()) + ". pop_10 is Steinsson's decadal copy of Clark with 1590 and "
    "1200 corrected (MalthusFigures Readme); decade value vs first-year value is not the same object")

# %% V09: BNS broad_pop_hf is the decade mean of A2 B
hf = S("bns_malthus_broad_pop_hf")
dm = b.groupby((b.index // 10) * 10).mean()
m9, n9, _ = maxrel(hf, dm)
rec("V09", "C1", "BNS broad_pop_hf vs decade mean of A2 B (Bank's 2024 file)", "PASS" if m9 < 1e-9 else "DIFF",
    f"{n9} decades, max rel {f4(m9)}. A decade mean of an interpolated series: before 1541 it is not data")
wp = S("bns_malthus_wrigley_pop")
m9b, n9b, _ = maxrel(wp, pop10)
rec("V09b", "C1", "BNS 'wrigley_pop'", "FLAG",
    f"= pop_10 (Clark) from 1540 ({n9b} decades, max rel {f4(m9b)}). It is Clark's column under Wrigley's name "
    "(dataset.R line 'wrigley_pop = ifelse(decade < 1540, NA, pop_10)'). Not a separate source")

# %% V10: what the BoE file carries from Clark, against the BNS copy of Clark's own files
af, bg, v = S("boe_A47_AF_clark_avg_male_wage"), S("boe_A47_BG_clark_col"), S("boe_A48_V_clark_real_earnings")
mv, nv, _ = maxrel(v, af / bg)
rec("V10a", "C1,C2", "A48 V = A47 AF / A47 BG", "PASS" if mv < 1e-12 else "DIFF", f"{nv} yrs, max rel {f4(mv)}")
for sid, c, lab in [("boe_A47_AF_clark_avg_male_wage", "G", "Male average Wage"),
                    ("boe_A47_BG_clark_col", "S", "Price Index - Cost of Living")]:
    m, n, d = maxrel(S(sid), col[c])
    rec("V10b", "C1,C2", f"{sid} vs Clark NNI 2015 col {c} '{lab}'", "PASS" if m < 1e-9 else "DIFF",
        f"{n} common yrs, max rel {f4(m)}" + (f"; differing years {list(d[d > 1e-9].index[:10])}" if m >= 1e-9 else ""))
m, n, d = maxrel(v, col["G"] / col["S"])
rec("V10c", "C1,C2", "A48 V vs Clark NNI 2015 G/S (the construction family B names 'Clark average wage / Clark COL')",
    "PASS" if m < 1e-9 else "DIFF", f"{n} common yrs, max rel {f4(m)}")
wbw0 = openpyxl.load_workbook(RAW / "bns" / "Wages 2014.xlsx", read_only=True, data_only=True)["Annual"]
wbw0.reset_dimensions()
w1848 = [r[6] for r in wbw0.iter_rows(min_row=2, max_row=671, values_only=True) if r[0] == 1848][0]
rec("V10g", "C1,C2", "Clark's own two files disagree in one year", "DIFF",
    f"cost of living 1848: NNI 2015 col S {col['S'][1848]:.4f}, Wages 2014 col G {w1848:.4f}; BoE A47 BG carries "
    f"{bg[1848]:.4f} (the Wages 2014 value). So A48 V differs from NNI G/S in 1848 only. Family B should see the same")
red = pd.read_csv(OUT / "boe" / "boe_A47_AF_clark_avg_male_wage.csv")
redy = list(red.loc[red["basis"].str.startswith("interpolated"), "year"])
nan_g = [y for y in redy if math.isnan(col["G"].get(y, float("nan")))]
rec("V10d", "C1", "A47 AF and BG: cells the Bank marks red (interpolated)", "FLAG",
    f"{len(redy)} years: {redy}. In Clark's NNI 2015 col G these years are "
    + ("blank" if len(nan_g) == len(redy) else f"present ({len(redy) - len(nan_g)} of {len(redy)})")
    + ": the Bank filled them. A48 V (AF/BG) is left blank in exactly these years: "
    + str(sorted(set(range(1209, 1870)) - set(v.index)) == sorted(redy)))
wbw = openpyxl.load_workbook(RAW / "bns" / "Wages 2014.xlsx", read_only=True, data_only=True)
ws = wbw["Annual"]
ws.reset_dimensions()
wrow = list(ws.iter_rows(min_row=2, max_row=671, values_only=True))
wy = [int(r[0]) for r in wrow]
wcol = {c: pd.Series([cell(r, c) for r in wrow], index=wy) for c in "BDEG"}
for sid, c, lab in [("boe_A47_Z_clark_farm_wage", "B", "Farm Laborers"), ("boe_A47_AC_clark_labourer_wage", "D", "Building Laborers"),
                    ("boe_A47_AB_clark_craftsman_wage", "E", "Building Craftsmen")]:
    s = S(sid)
    m, n, d = maxrel(s[s.index <= 1869], wcol[c])
    res = "PASS" if m < 1e-9 else "DIFF"
    extra = f"; {int((d > 1e-9).sum())} years differ, e.g. {list(d[d > 1e-9].index[:8])}" if m >= 1e-9 else ""
    rec("V10e", "C1,C2", f"{sid} (BoE says Clark 2009) vs Clark Wages 2014.xlsx Annual col {c} '{lab}'", res,
        f"{n} common yrs 1209-1869, max rel {f4(m)}{extra}" + (f"; BoE runs on to {s.index.max()}" if s.index.max() > 1869 else ""))
m, n, d = maxrel(bg, wcol["G"])
rec("V10f", "C1", "A47 BG vs Wages 2014 Annual col G 'Cost of Living'", "PASS" if m < 1e-9 else "DIFF",
    f"{n} common yrs, max rel {f4(m)}" + (f"; e.g. {list(d[d > 1e-9].index[:8])}" if m >= 1e-9 else ""))

# %% V11: Humphries-Weisdorf, the 2016 version in BoE against the 2019 version in BNS
e = S("boe_A48_E_hw2016_real_earnings")
bh = S("boe_A47_BH_allen_cpi_respectable")
hwf = S("bns_hw19_F_real_income_annual")
eb = (e * bh.reindex(e.index)).dropna()
# decades as in HW: 1260-70, 1270-80, ...: take the BoE decade mean over years d..d+9
em = e.groupby((e.index // 10) * 10).mean()
cmp_ = {int(d): (em.get(d, float("nan")), hwf[d]) for d in hwf.index}
rat = pd.Series({d: a / b_ for d, (a, b_) in cmp_.items() if not math.isnan(a)})
rec("V11a", "C1", "BoE A48 E (HW 2016 working paper) vs BNS hw19 col F (HW 2019, EJ), by decade", "DIFF",
    f"{len(rat)} decades; ratio BoE-decade-mean / HW2019: min {rat.min():.3f} ({rat.idxmin()}), max {rat.max():.3f} "
    f"({rat.idxmax()}), median {rat.median():.3f}. Two versions of one construction; report both. Sample: "
    + ", ".join(f"{d}:{rat[d]:.3f}" for d in [1260, 1340, 1350, 1400, 1500, 1600, 1700, 1750, 1800, 1840] if d in rat.index))
grp = (eb.index - 1) // 10  # decades 1261-1270, ..., 1701-1710: where the product steps
dev = (eb / eb.groupby(grp).transform("median") - 1).abs().groupby(grp).max()
lvl = eb.groupby(grp).median()
step = (lvl / lvl.shift(1) - 1).abs().dropna()
two_dp = bool(((e * 100).round() - e * 100).abs().max() < 1e-9)
rec("V11b", "C1", "BoE A48 E: how it is built", "FLAG",
    f"E is typed to 2 decimals: {two_dp}. E x Allen CPI (A47 BH) moves in steps at 1261, 1271, ..., 1401, 1411, ...: "
    f"median step between decades {100 * step.median():.1f}%, median within-decade spread {100 * dev.median():.1f}% "
    f"(max {100 * dev.max():.1f}%). The numerator is decadal; the denominator is an annual price close to, not "
    "exactly, A47 BH (the BoE note adds a 5% rent adjustment). Within-decade movement is prices. 1265-1850")

# %% V12: Allen (2007) prices, BoE vs BNS digitisation
m, n, d = maxrel(S("boe_A47_BL_allen2007_col"), S("bns_allen07_cpi"))
rec("V12", "C2", "BoE A47 BL vs BNS allen07_app1 (Allen 2007 COL)", "PASS" if m < 1e-9 else "DIFF",
    f"{n} yrs 1770-1869, max rel {f4(m)}" + (f"; {list(d[d > 1e-9].index[:10])}" if m >= 1e-9 else ""))

# %% V13: the Feinstein band: one wage, two deflators
w_, x_, y_, z_ = (S(f"boe_A48_{c}") for c in ["W_feinstein_real", "X_feinstein_wage_allen_prices",
                                              "Y_feinstein_uk_wage_gb_prices", "Z_feinstein_uk_wage_allen_prices"])
r = (x_ / w_).dropna()
rec("V13a", "C2", "A48 X/W: Feinstein GB wages deflated by Allen (2007) vs Feinstein (1998) prices, both 1835=100", "DIFF",
    ", ".join(f"{yy}:{r[yy]:.3f}" for yy in [1770, 1780, 1800, 1815, 1830, 1850, 1869]) +
    f". Growth 1770->1869: W x{w_[1869] / w_[1770]:.2f}, X x{x_[1869] / x_[1770]:.2f}. A band, not a choice")
r2 = (z_ / y_).dropna()
rec("V13b", "C2", "A48 Z/Y: same with UK wages", "DIFF", ", ".join(f"{yy}:{r2[yy]:.3f}" for yy in [1770, 1800, 1830, 1850, 1869]))
aq, bk, bl = S("boe_A47_AQ_feinstein_earnings_gbp"), S("boe_A47_BK_feinstein_col"), S("boe_A47_BL_allen2007_col")
ar = S("boe_A47_AR_feinstein_earnings_index")
mw, nw_, _ = maxrel(w_, 100 * (ar / bk) / (ar[1835] / bk[1835]))
mx_, nx_, _ = maxrel(x_, 100 * (aq / bl) / (aq[1835] / bl[1835]))
rec("V13c", "C2", "A48 W and X recomputed from A47 (AR/BK and AQ/BL, rebased 1835)", "PASS" if max(mw, mx_) < 1e-9 else "DIFF",
    f"W {nw_} yrs max rel {f4(mw)}; X {nx_} yrs max rel {f4(mx_)}")

# %% V14: the BoE composite real wage (A48 B) is a splice
bb_, b47, d47 = S("boe_A48_B_real_earnings_composite"), S("boe_A47_B_earnings_composite"), S("boe_A47_D_cpi_composite")
m, n, _ = maxrel(bb_, 100 * (b47 / d47) / (b47[1900] / d47[1900]))
rec("V14a", "C1,C2", "A48 B = 100 x (A47 B / A47 D) / (1900 value)", "PASS" if m < 1e-9 else "DIFF", f"{n} yrs, max rel {f4(m)}")
rec("V14b", "C1,C2", "A48 B links (read from A47 formulas)", "FLAG",
    "wage growth: 1209-1750 Clark avg male daily wage (England, A47 AF); 1750-1770 Crafts-Mills weekly earnings (GB, AP); "
    "1770-1881 Feinstein 1998 (GB, AR); 1881-1911 Feinstein 1990 (AS); 1911-1913 BC; 1913-1963 Feinstein UK (BD); "
    "1963-2016 ONS (BF). Price growth: 1209-1661 Clark COL (England, BG); 1661-1750 Schumpeter-Gilboy (England, BI); "
    "1750-1770 Crafts-Mills (GB, BJ); 1770-1882 Feinstein 1998 (GB, BK); 1882-1914 Feinstein 1991 (BW); 1914-2016 "
    "ONS splice (F). Chain-linked backwards from the modern level. Header says 'GB'; it is England before 1750. "
    "Name it 'BoE composite real consumption earnings'")
bv = (bb_ / v).dropna()
rec("V14c", "C1", "A48 B / A48 V (BoE composite vs Clark real earnings), 1209-1869", "DIFF",
    f"constant to 1661 (both on Clark wage and Clark COL): 1300 {bv[1300]:.4f}, 1660 {bv[1660]:.4f}. Then it drifts: "
    f"1700 {bv[1700]:.4f}, 1749 {bv[1749]:.4f} (Schumpeter-Gilboy vs Clark COL), 1769 {bv[1769]:.4f}, 1800 {bv[1800]:.4f}, "
    f"1830 {bv[1830]:.4f}, 1850 {bv[1850]:.4f}, 1869 {bv[1869]:.4f}. Relative to 1660: 1869 is "
    f"{100 * (bv[1869] / bv[1660] - 1):+.0f}%. Two constructions over the same years; not reconciled")

# %% V15-V16: Clark's NNI in the BoE memo column, and England GDP against it
cw = S("boe_A9_CW_clark_nni")
m, n, d = maxrel(cw, col["O"])
rec("V15", "C3", "BoE A9 CW (memo 'Clark (2015)' NNI) vs Clark NNI 2015 col O", "PASS" if m < 1e-9 else "DIFF",
    f"{n} yrs 1700-1869, max rel {f4(m)}")
gd = S("boe_A21_D_gdp_england_nominal_fc")
rr = (gd / cw).dropna()
rec("V16", "C3", "BoE England nominal GDP at factor cost (A21 D) / Clark NNI (A9 CW)", "DIFF",
    ", ".join(f"{yy}:{rr[yy]:.3f}" for yy in [1700, 1750, 1770, 1800, 1820, 1850, 1869]) +
    ". Different objects (GDP vs NNI) and different builders. A land share over the BoE denominator is not Clark's share")
cu = S("boe_A9_CU_england_share_of_gb")
rec("V16b", "C3", "BoE England nominal GDP 1700-1910 is a construction", "FLAG",
    f"= England share of GB (A9 CU, 'Geary and Stark (2015a) and Broadberry et al. (2015), interpolated': {cu[1700]:.1f}% "
    f"in 1700, {cu[1800]:.1f}% in 1800, {cu[1900]:.1f}% in 1900) x GB GDP at factor cost (A9 BM). Before 1700 chained on "
    "Broadberry A6 col G")

# %% V17: rent after 1855: land and buildings, UK: an upper bound only
am, arr, aw = S("boe_A17_AM_rent_mitchell"), S("boe_A17_AR_gdpi_mitchell_fc"), S("boe_A17_AW_gdpi_composite")
sh = (am / arr).dropna()
sw = (am / aw).dropna()
clark_share = (col["J"] / col["O"]).dropna()
rec("V17a", "C3", "Mitchell rent / Mitchell GDP(I) at factor cost (A17 AM / AR), UK, land AND buildings", "FLAG",
    ", ".join(f"{yy}:{sh[yy]:.3f}" for yy in [1855, 1860, 1869, 1880, 1890, 1900, 1913, 1920]) +
    ". Our named construction: an UPPER BOUND on a land share, not a land share. UK, not England. The fall "
    f"1913->1920 ({sh[1913]:.3f} -> {sh[1920]:.3f}: AM {am[1913]:.0f} -> {am[1920]:.0f}, AR {arr[1913]:.0f} -> "
    f"{arr[1920]:.0f}) is not examined here")
rec("V17b", "C3", "same over the BoE composite GDP(I) (AM / AW)", "INFO",
    ", ".join(f"{yy}:{sw[yy]:.3f}" for yy in [1855, 1869, 1900, 1913, 1920]) +
    f". AW is chained on AR to 1919 with its level set later: AW/AR = {aw[1855] / arr[1855]:.4f} (1855), "
    f"{aw[1913] / arr[1913]:.4f} (1913). Prefer AM/AR (one table)")
rec("V17c", "C3", "overlap 1855-1869: Clark land rents / NNI (England, land only; BNS copy of Clark) beside AM/AR", "DIFF",
    ", ".join(f"{yy}: Clark {clark_share[yy]:.3f} vs AM/AR {sh[yy]:.3f}" for yy in [1855, 1860, 1865, 1869]) +
    ". Different objects and geographies. Not spliced")
rec("V17d", "C3", "BoE land series", "INFO",
    "the file carries no agricultural land rent. Land appears only as: A3 B arable and sown acreage (England 1270-1871, "
    "Broadberry); A32 AC agricultural land PRICE (GBP 000 per acre, 1892-1969, Vallis 1972); A21 C a factor-cost "
    "adjustment that treats direct taxes on land as a tax on production (1270-1700)")

# %% V18: BNS malthus_data columns against their stated inputs
checks = [("bns_malthus_land_share", "bns_clark10_t13_land_share", "land_share = Clark (2010) Table 13 Land"),
          ("bns_malthus_inc_hw", "bns_hw19_F_real_income_annual", "inc_hw = HW (2019) Table A2 col F")]
for a_, b2, lab in checks:
    m, n, _ = maxrel(S(a_), S(b2))
    rec("V18", "C1", lab, "PASS" if m < 1e-9 else "DIFF", f"{n} decades, max rel {f4(m)}")
aw_ = S("bns_malthus_allen_w")
wl = S("bns_malthus_w_lab_10")
a07 = S("bns_allen07_wage")
a07d = a07.groupby((a07.index // 10) * 10).mean()
norm = wl[1770] / a07d[1770]
inside = aw_[(aw_.index >= 1770) & (aw_.index <= 1860)]
outside = aw_[(aw_.index < 1770) | (aw_.index > 1860)]
m_in, n_in, _ = maxrel(inside, a07d * norm)
m_out, n_out, _ = maxrel(outside, wl)
rec("V18", "C1,C2", "BNS allen_w", "FLAG",
    f"1770-1860 = Allen (2007) wage decade mean x {norm:.4f} (Clark w_lab_10 / Allen in 1770s): {n_in} decades, max rel "
    f"{f4(m_in)}. Before 1770 it IS Clark's real building labourer wage ({n_out} decades, max rel {f4(m_out)}). "
    "A splice under Allen's name; use only 1770-1860 as Allen")
lf = S("bns_malthus_broad_pop_lf")
rec("V18", "C1", "BNS broad_pop_lf", "FLAG",
    f"Broadberry benchmarks placed in their decade before 1540 ({int((lf.index < 1540).sum())} decades), broad_pop_hf "
    "after: benchmarks without interpolation, then Wrigley-era annual means")

# %% V19: descriptive numbers for the eyeball sheet (no fit)
pdm = b.groupby((b.index // 10) * 10).mean()
vdm = v.groupby((v.index // 10) * 10).mean()
edm = e.groupby((e.index // 10) * 10).mean()
rec("V19a", "C1", "decade means: population (A2 B, m) | Clark real earnings (A48 V) | HW2016 (A48 E)", "INFO",
    "; ".join(f"{d}s {pdm[d]:.2f} | {vdm[d]:.3f} | {f4(edm.get(d, float('nan')))}" for d in
              [1290, 1310, 1340, 1350, 1360, 1380, 1400, 1450, 1500, 1550, 1600, 1650, 1700, 1750, 1800]))
rec("V19b", "C2", "BoE composite real earnings (A48 B, 1900=100)", "INFO",
    ", ".join(f"{yy}:{bb_[yy]:.1f}" for yy in [1700, 1750, 1770, 1800, 1820, 1840, 1850, 1860, 1870, 1880, 1900, 1914, 1950]))

# %% write and print
with open(OUT / "validation_family_a.tsv", "w", encoding="utf-8", newline="") as f:
    wr = csv.DictWriter(f, fieldnames=["check", "criterion", "subject", "result", "detail"], delimiter="\t", lineterminator="\n")
    wr.writeheader()
    wr.writerows(results)
for r_ in results:
    print(f"{r_['check']:5} {r_['result']:4} [{r_['criterion']}] {r_['subject']}\n      {r_['detail']}")
print("STOP" if any(r_["result"] == "STOP" for r_ in results) else "no STOP", "-", len(results), "checks ->",
      OUT / "validation_family_a.tsv")
